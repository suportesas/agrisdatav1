"""
Utilitários para geração de relatório de estoque por produto
"""
import io
from datetime import datetime
from decimal import Decimal
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.db.models import Sum, Q
from estoque.models import Lote, MovimentacaoEstoque


class RelatorioProdutoPDF:
    """Gerador de relatório de estoque por produto em PDF"""
    
    def __init__(self, empresa, titulo='Relatório de Estoque por Produto'):
        self.empresa = empresa
        self.titulo = titulo
        self.buffer = io.BytesIO()
        self.doc = SimpleDocTemplate(
            self.buffer,
            pagesize=A4,
            rightMargin=0.5*inch,
            leftMargin=0.5*inch,
            topMargin=0.75*inch,
            bottomMargin=0.5*inch,
        )
        self.styles = getSampleStyleSheet()
        self.story = []
        self._setup_styles()
    
    def _setup_styles(self):
        """Configurar estilos customizados"""
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=6,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold'
        ))
        
        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#666666'),
            spaceAfter=12,
            alignment=TA_CENTER,
        ))
    
    def adicionar_cabecalho(self, filtros_desc=''):
        """Adicionar cabeçalho ao relatório"""
        self.story.append(Paragraph(self.titulo, self.styles['CustomTitle']))
        
        info_texto = f'<b>Empresa:</b> {self.empresa.nome}<br/>'
        info_texto += f'<b>Gerado em:</b> {datetime.now().strftime("%d/%m/%Y às %H:%M")}<br/>'
        if filtros_desc:
            info_texto += f'<b>Filtros:</b> {filtros_desc}'
        
        self.story.append(Paragraph(info_texto, self.styles['Subtitle']))
        self.story.append(Spacer(1, 0.3*inch))
    
    def adicionar_tabela_produtos(self, produtos_data):
        """Adicionar tabela de produtos ao relatório"""
        if not produtos_data:
            self.story.append(Paragraph('Nenhum produto encontrado com os critérios especificados.', self.styles['Normal']))
            return
        
        # Preparar dados
        dados = [['Produto', 'Categoria', 'Est. Atual', 'Est. Mínimo', 'Entrada', 'Saída', 'Status']]
        
        for item in produtos_data:
            dados.append([
                item['produto'][:35],
                item['categoria'][:20],
                item['estoque_atual'],
                item['estoque_minimo'],
                item['entrada'],
                item['saida'],
                item['status'],
            ])
        
        # Criar tabela
        tabela = Table(dados, colWidths=[2*inch, 1.2*inch, 0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch, 0.8*inch])
        
        # Aplicar estilos
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        
        self.story.append(tabela)
    
    def gerar(self):
        """Gerar o PDF"""
        self.doc.build(self.story)
        self.buffer.seek(0)
        return self.buffer.getvalue()


class RelatorioProdutoExcel:
    """Gerador de relatório de estoque por produto em Excel"""
    
    def __init__(self, empresa, titulo='Relatório de Estoque por Produto'):
        self.empresa = empresa
        self.titulo = titulo
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Estoque por Produto'
        self._setup_styles()
    
    def _setup_styles(self):
        """Configurar estilos do Excel"""
        self.header_fill = PatternFill(start_color='1f4788', end_color='1f4788', fill_type='solid')
        self.header_font = Font(bold=True, color='FFFFFF', size=11)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def adicionar_cabecalho(self, filtros_desc=''):
        """Adicionar cabeçalho"""
        self.ws['A1'] = self.titulo
        self.ws['A1'].font = Font(bold=True, size=14, color='1f4788')
        
        self.ws['A2'] = f'Empresa: {self.empresa.nome}'
        self.ws['A3'] = f'Gerado em: {datetime.now().strftime("%d/%m/%Y às %H:%M")}'
        
        if filtros_desc:
            self.ws['A4'] = f'Filtros: {filtros_desc}'
        
        return 5
    
    def adicionar_tabela_produtos(self, produtos_data, linha_inicio=5):
        """Adicionar tabela de produtos"""
        # Cabeçalhos
        colunas = ['Produto', 'Categoria', 'Est. Atual', 'Unid.', 'Est. Mínimo', 'Entrada Total', 'Saída Total', 'Status']
        
        for col, titulo in enumerate(colunas, 1):
            cell = self.ws.cell(row=linha_inicio, column=col)
            cell.value = titulo
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Dados
        linha = linha_inicio + 1
        for item in produtos_data:
            self.ws.cell(row=linha, column=1).value = item['produto']
            self.ws.cell(row=linha, column=2).value = item['categoria']
            self.ws.cell(row=linha, column=3).value = item['estoque_atual_num']
            self.ws.cell(row=linha, column=4).value = item['unidade']
            self.ws.cell(row=linha, column=5).value = item['estoque_minimo_num']
            self.ws.cell(row=linha, column=6).value = item['entrada_num']
            self.ws.cell(row=linha, column=7).value = item['saida_num']
            self.ws.cell(row=linha, column=8).value = item['status']
            
            # Aplicar estilos
            for col in range(1, 9):
                cell = self.ws.cell(row=linha, column=col)
                cell.border = self.border
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
                # Alinhar números à direita
                if col in [3, 5, 6, 7]:
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    cell.number_format = '0.00'
            
            linha += 1
        
        # Ajustar largura das colunas
        self.ws.column_dimensions['A'].width = 35
        self.ws.column_dimensions['B'].width = 20
        self.ws.column_dimensions['C'].width = 12
        self.ws.column_dimensions['D'].width = 8
        self.ws.column_dimensions['E'].width = 12
        self.ws.column_dimensions['F'].width = 15
        self.ws.column_dimensions['G'].width = 15
        self.ws.column_dimensions['H'].width = 15
        
        return linha
    
    def gerar(self):
        """Gerar o Excel"""
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


def gerar_dados_relatorio_produto(lotes, empresa, incluir_zerados=False, ordenacao='produto'):
    """
    Gera dados agregados por produto para o relatório
    
    Args:
        lotes: QuerySet de lotes filtrados
        empresa: Empresa para filtrar movimentações
        incluir_zerados: Se deve incluir produtos com estoque zerado
        ordenacao: Campo para ordenação ('produto', 'categoria', 'estoque', 'status')
    """
    from produtos.models import Produto
    
    # Agregar lotes por produto
    produtos_agregados = lotes.values(
        'produto_id',
        'produto__nome',
        'produto__categoria__nome',
        'produto__estoque_atual',
        'produto__estoque_minimo',
        'unidade_medida__sigla'
    ).annotate(
        total_estoque=Sum('quantidade_atual')
    )
    
    produto_ids = [p['produto_id'] for p in produtos_agregados]
    
    # Buscar movimentações de entrada e saída
    entradas_q = Q(tipo='ENTRADA') | Q(origem__in=['AJUSTE_ENTRADA', 'TRANSFERENCIA_ENTRADA', 'DEVOLUCAO_CLIENTE', 'PRODUCAO', 'COMPRA'])
    saidas_q = Q(tipo='SAIDA') | Q(origem__in=['AJUSTE_SAIDA', 'TRANSFERENCIA_SAIDA', 'DEVOLUCAO_FORNECEDOR', 'REQUISICAO', 'CONSUMO_PRODUCAO', 'VENDA', 'PERDA'])
    
    entradas = MovimentacaoEstoque.objects.filter(
        empresa=empresa,
        status='EXECUTADO',
        lote__produto_id__in=produto_ids
    ).filter(entradas_q).values('lote__produto_id').annotate(total=Sum('quantidade'))
    
    saidas = MovimentacaoEstoque.objects.filter(
        empresa=empresa,
        status='EXECUTADO',
        lote__produto_id__in=produto_ids
    ).filter(saidas_q).values('lote__produto_id').annotate(total=Sum('quantidade'))
    
    entradas_map = {e['lote__produto_id']: e['total'] for e in entradas}
    saidas_map = {s['lote__produto_id']: s['total'] for s in saidas}
    
    # Montar dados finais
    dados = []
    for item in produtos_agregados:
        estoque_atual = float(item['produto__estoque_atual'] or 0)
        estoque_minimo = float(item['produto__estoque_minimo'] or 0)
        
        # Pular produtos zerados se não incluir
        if not incluir_zerados and estoque_atual == 0:
            continue
        
        entrada_total = float(entradas_map.get(item['produto_id'], 0) or 0)
        saida_total = float(saidas_map.get(item['produto_id'], 0) or 0)
        unidade = item['unidade_medida__sigla'] or ''
        
        # Determinar status
        if estoque_atual == 0:
            status = 'Zerado'
        elif estoque_atual < estoque_minimo:
            status = 'Baixo'
        else:
            status = 'Normal'
        
        dados.append({
            'produto': item['produto__nome'],
            'categoria': item['produto__categoria__nome'] or 'Sem categoria',
            'estoque_atual': f"{estoque_atual:.2f} {unidade}",
            'estoque_atual_num': estoque_atual,
            'estoque_minimo': f"{estoque_minimo:.2f} {unidade}",
            'estoque_minimo_num': estoque_minimo,
            'entrada': f"{entrada_total:.2f} {unidade}",
            'entrada_num': entrada_total,
            'saida': f"{saida_total:.2f} {unidade}",
            'saida_num': saida_total,
            'unidade': unidade,
            'status': status,
        })
    
    # Ordenar conforme solicitado
    if ordenacao == 'categoria':
        dados.sort(key=lambda x: x['categoria'])
    elif ordenacao == 'estoque':
        dados.sort(key=lambda x: x['estoque_atual_num'], reverse=True)
    elif ordenacao == 'status':
        status_order = {'Zerado': 0, 'Baixo': 1, 'Normal': 2}
        dados.sort(key=lambda x: (status_order.get(x['status'], 99), x['produto']))
    else:  # Default: produto
        dados.sort(key=lambda x: x['produto'])
    
    return dados
