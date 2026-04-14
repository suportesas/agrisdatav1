import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer


def gerar_extrato_emprestimo_excel(empresa, emprestimos):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Extrato Empréstimos'

    ws['A1'] = 'Extrato de Empréstimos de Materiais'
    ws['A2'] = f'Empresa: {empresa.nome}'
    ws['A3'] = f'Gerado em: {datetime.now().strftime("%d/%m/%Y %H:%M")}'

    header_fill = PatternFill(start_color='1F4788', end_color='1F4788', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    headers = [
        'Data', 'Produtor', 'Produto', 'Status', 'Promessa',
        'Qtd Emprestada', 'Qtd Devolvida', 'Saldo', 'Requisição'
    ]

    for col_idx, title in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=col_idx, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    row = 6
    for item in emprestimos:
        ws.cell(row=row, column=1, value=item.data_emprestimo.strftime('%d/%m/%Y %H:%M') if item.data_emprestimo else '')
        ws.cell(row=row, column=2, value=item.produtor.razao_social)
        ws.cell(row=row, column=3, value=item.produto.nome)
        ws.cell(row=row, column=4, value=item.get_status_display())
        ws.cell(row=row, column=5, value=item.data_promessa_devolucao.strftime('%d/%m/%Y') if item.data_promessa_devolucao else '-')
        ws.cell(row=row, column=6, value=float(item.quantidade_emprestada or 0))
        ws.cell(row=row, column=7, value=float(item.quantidade_devolvida or 0))
        ws.cell(row=row, column=8, value=float(item.saldo_pendente or 0))
        ws.cell(row=row, column=9, value=item.requisicao.numero if item.requisicao else '-')
        row += 1

    widths = [18, 34, 30, 22, 15, 16, 16, 12, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def gerar_extrato_emprestimo_pdf(empresa, emprestimos):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=0.4 * inch,
        leftMargin=0.4 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.4 * inch,
    )

    styles = getSampleStyleSheet()
    story = [
        Paragraph('Extrato de Empréstimos de Materiais', styles['Heading2']),
        Paragraph(f'Empresa: {empresa.nome}', styles['Normal']),
        Paragraph(f'Gerado em: {datetime.now().strftime("%d/%m/%Y %H:%M")}', styles['Normal']),
        Spacer(1, 0.2 * inch),
    ]

    data = [[
        'Data', 'Produtor', 'Produto', 'Status', 'Promessa',
        'Emprestado', 'Devolvido', 'Saldo'
    ]]

    for item in emprestimos:
        data.append([
            item.data_emprestimo.strftime('%d/%m/%Y') if item.data_emprestimo else '',
            item.produtor.razao_social[:24],
            item.produto.nome[:24],
            item.get_status_display(),
            item.data_promessa_devolucao.strftime('%d/%m/%Y') if item.data_promessa_devolucao else '-',
            f'{item.quantidade_emprestada:.3f}',
            f'{item.quantidade_devolvida:.3f}',
            f'{item.saldo_pendente:.3f}',
        ])

    table = Table(data, colWidths=[0.9 * inch, 1.4 * inch, 1.4 * inch, 0.9 * inch, 0.85 * inch, 0.75 * inch, 0.75 * inch, 0.65 * inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#999999')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f6f6f6')]),
        ('ALIGN', (5, 1), (-1, -1), 'RIGHT'),
    ]))

    story.append(table)
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
