"""
Utilitários para geração de relatório de pedidos pendentes
"""
import io
from datetime import datetime
from decimal import Decimal
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class RelatorioPedidosPendentesPDF:
    """Gerador de relatório de pedidos pendentes em PDF"""
    
    def __init__(self, empresa, titulo='Relatório de Pedidos Pendentes'):
        self.empresa = empresa
        self.titulo = titulo
        self.buffer = io.BytesIO()
        self.doc = SimpleDocTemplate(
            self.buffer,
            pagesize=A4,
            rightMargin=0.5*inch,
            leftMargin=0.5*inch,
            topMargin=0.75*inch,
            bottomMargin=0.5*inch,
        )
        self.styles = getSampleStyleSheet()
        self.story = []
        self._setup_styles()
    
    def _setup_styles(self):
        """Configurar estilos customizados"""
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=6,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold'
        ))
        
        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#666666'),
            spaceAfter=12,
            alignment=TA_CENTER,
        ))
    
    def adicionar_cabecalho(self, filtros_desc=''):
        """Adicionar cabeçalho ao relatório"""
        self.story.append(Paragraph(self.titulo, self.styles['CustomTitle']))
        
        info_texto = f'<b>Empresa:</b> {self.empresa.nome}<br/>'
        info_texto += f'<b>Gerado em:</b> {datetime.now().strftime("%d/%m/%Y às %H:%M")}<br/>'
        if filtros_desc:
            info_texto += f'<b>Filtros:</b> {filtros_desc}'
        
        self.story.append(Paragraph(info_texto, self.styles['Subtitle']))
        self.story.append(Spacer(1, 0.3*inch))
    
    def adicionar_tabela_pedidos(self, pedidos_data):
        """Adicionar tabela de pedidos ao relatório"""
        if not pedidos_data:
            self.story.append(Paragraph('Nenhum pedido pendente encontrado.', self.styles['Normal']))
            return
        
        # Preparar dados
        dados = [['Pedido', 'Data', 'Fornecedor', 'Valor Total', 'Qtd. Recebida', 'Qtd. Pendente', 'Status']]
        
        for item in pedidos_data:
            dados.append([
                item['numero'],
                item['data'],
                item['fornecedor'][:25],
                item['valor_total'],
                item['qtd_recebida'],
                item['qtd_pendente'],
                item['status'],
            ])
        
        # Criar tabela
        tabela = Table(dados, colWidths=[1*inch, 0.9*inch, 2*inch, 1*inch, 1*inch, 1*inch, 1*inch])
        
        # Aplicar estilos
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#d9534f')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        
        self.story.append(tabela)
    
    def gerar(self):
        """Gerar o PDF"""
        self.doc.build(self.story)
        self.buffer.seek(0)
        return self.buffer.getvalue()


class RelatorioPedidosPendentesExcel:
    """Gerador de relatório de pedidos pendentes em Excel"""
    
    def __init__(self, empresa, titulo='Relatório de Pedidos Pendentes'):
        self.empresa = empresa
        self.titulo = titulo
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Pedidos Pendentes'
        self._setup_styles()
    
    def _setup_styles(self):
        """Configurar estilos do Excel"""
        self.header_fill = PatternFill(start_color='d9534f', end_color='d9534f', fill_type='solid')
        self.header_font = Font(bold=True, color='FFFFFF', size=11)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def adicionar_cabecalho(self, filtros_desc=''):
        """Adicionar cabeçalho"""
        self.ws['A1'] = self.titulo
        self.ws['A1'].font = Font(bold=True, size=14, color='d9534f')
        
        self.ws['A2'] = f'Empresa: {self.empresa.nome}'
        self.ws['A3'] = f'Gerado em: {datetime.now().strftime("%d/%m/%Y às %H:%M")}'
        
        if filtros_desc:
            self.ws['A4'] = f'Filtros: {filtros_desc}'
        
        return 5
    
    def adicionar_tabela_pedidos(self, pedidos_data, linha_inicio=5):
        """Adicionar tabela de pedidos"""
        # Cabeçalhos
        colunas = ['Pedido', 'Data', 'Fornecedor', 'Valor Total', 'Qtd. Recebida', 'Qtd. Pendente', 'Status']
        
        for col, titulo in enumerate(colunas, 1):
            cell = self.ws.cell(row=linha_inicio, column=col)
            cell.value = titulo
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Dados
        linha = linha_inicio + 1
        for item in pedidos_data:
            self.ws.cell(row=linha, column=1).value = item['numero']
            self.ws.cell(row=linha, column=2).value = item['data']
            self.ws.cell(row=linha, column=3).value = item['fornecedor']
            self.ws.cell(row=linha, column=4).value = item['valor_total_num']
            self.ws.cell(row=linha, column=5).value = item['qtd_recebida_num']
            self.ws.cell(row=linha, column=6).value = item['qtd_pendente_num']
            self.ws.cell(row=linha, column=7).value = item['status']
            
            # Aplicar estilos
            for col in range(1, 8):
                cell = self.ws.cell(row=linha, column=col)
                cell.border = self.border
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
                # Alinhar números à direita
                if col in [4, 5, 6]:
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    if col == 4:
                        cell.number_format = 'R$ #,##0.00'
                    else:
                        cell.number_format = '0.00'
            
            linha += 1
        
        # Ajustar largura das colunas
        self.ws.column_dimensions['A'].width = 15
        self.ws.column_dimensions['B'].width = 12
        self.ws.column_dimensions['C'].width = 30
        self.ws.column_dimensions['D'].width = 15
        self.ws.column_dimensions['E'].width = 15
        self.ws.column_dimensions['F'].width = 15
        self.ws.column_dimensions['G'].width = 20
        
        return linha
    
    def gerar(self):
        """Gerar o Excel"""
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


def gerar_dados_pedidos_pendentes(pedidos):
    """
    Gera dados de pedidos pendentes para o relatório
    """
    dados = []
    
    for pedido in pedidos:
        # Calcular quantidades
        total_pedido = sum(item.quantidade_pedida for item in pedido.itens.all())
        # Usa FINALIZADO pois é o status quando o recebimento foi concluído
        total_recebido = sum(
            item.quantidade_recebida 
            for receb in pedido.recebimentos.filter(status='FINALIZADO') 
            for item in receb.itens.all()
        )
        total_pendente = total_pedido - total_recebido
        
        # Determinar status
        if total_recebido == 0:
            status = 'Não Recebido'
        elif total_recebido < total_pedido:
            status = 'Parcial'
        else:
            status = 'Completo'
        
        dados.append({
            'numero': pedido.numero,
            'data': pedido.data_pedido.strftime('%d/%m/%Y'),
            'fornecedor': pedido.fornecedor.nome,
            'valor_total': f'R$ {float(pedido.valor_total):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.'),
            'valor_total_num': float(pedido.valor_total),
            'qtd_recebida': f'{total_recebido:.2f}',
            'qtd_recebida_num': total_recebido,
            'qtd_pendente': f'{total_pendente:.2f}',
            'qtd_pendente_num': total_pendente,
            'status': status,
        })
    
    return dados
