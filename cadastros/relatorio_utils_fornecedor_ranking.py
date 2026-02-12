"""
Utilitarios para geracao de relatorios de ranking de compras por fornecedor
"""
import io
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.utils import ImageReader
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def _create_logo_cell(logo_path, max_width, max_height):
    try:
        reader = ImageReader(logo_path)
        img_width, img_height = reader.getSize()
        if not img_width or not img_height:
            return ''
        scale = min(max_width / img_width, max_height / img_height)
        return Image(logo_path, width=img_width * scale, height=img_height * scale)
    except Exception:
        return ''


class RelatorioRankingFornecedoresPDF:
    def __init__(self, empresa, titulo='Ranking de Compras por Fornecedor'):
        self.empresa = empresa
        self.titulo = titulo
        self.buffer = io.BytesIO()
        self.doc = SimpleDocTemplate(
            self.buffer,
            pagesize=A4,
            rightMargin=0.5 * inch,
            leftMargin=0.5 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.5 * inch,
        )
        self.styles = getSampleStyleSheet()
        self.story = []
        self._setup_styles()

    def _setup_styles(self):
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=6,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold'
        ))

        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#666666'),
            spaceAfter=12,
            alignment=TA_CENTER,
        ))

    def adicionar_cabecalho(self, filtros_desc=''):
        info_texto = f'<b>Empresa:</b> {self.empresa.nome}<br/>'
        info_texto += f'<b>Gerado em:</b> {datetime.now().strftime("%d/%m/%Y as %H:%M")}<br/>'
        if filtros_desc:
            info_texto += f'<b>Filtros:</b> {filtros_desc}'

        logo_cell = ''
        try:
            if getattr(self.empresa, 'logo', None) and getattr(self.empresa.logo, 'path', None):
                logo_cell = _create_logo_cell(self.empresa.logo.path, 1.6 * inch, 1.0 * inch)
        except Exception:
            logo_cell = ''

        header_table = Table(
            [[logo_cell, Paragraph(info_texto, self.styles['Subtitle'])]],
            colWidths=[1.8 * inch, 4.8 * inch]
        )
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))

        self.story.append(header_table)
        self.story.append(Spacer(1, 0.12 * inch))
        self.story.append(Paragraph(self.titulo, self.styles['CustomTitle']))
        self.story.append(Spacer(1, 0.3 * inch))

    def adicionar_tabela(self, ranking):
        if not ranking:
            self.story.append(Paragraph('Nenhum registro encontrado com os criterios especificados.', self.styles['Normal']))
            return

        dados = [['Fornecedor', 'Total Comprado', '% do Total']]

        for item in ranking:
            dados.append([
                (item['fornecedor'] or '')[:60],
                item['total_formatado'],
                item['percentual_formatado'],
            ])

        tabela = Table(dados, colWidths=[3.4 * inch, 1.6 * inch, 1.0 * inch])
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))

        self.story.append(tabela)

    def gerar(self):
        self.doc.build(self.story)
        self.buffer.seek(0)
        return self.buffer.getvalue()


class RelatorioRankingFornecedoresExcel:
    def __init__(self, empresa, titulo='Ranking de Compras por Fornecedor'):
        self.empresa = empresa
        self.titulo = titulo
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Ranking'
        self._setup_styles()

    def _setup_styles(self):
        self.header_fill = PatternFill(start_color='1f4788', end_color='1f4788', fill_type='solid')
        self.header_font = Font(bold=True, color='FFFFFF', size=11)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def adicionar_cabecalho(self, filtros_desc=''):
        self.ws['A1'] = self.titulo
        self.ws['A1'].font = Font(bold=True, size=14, color='1f4788')

        self.ws['A2'] = f'Empresa: {self.empresa.nome}'
        self.ws['A3'] = f'Gerado em: {datetime.now().strftime("%d/%m/%Y as %H:%M")}'

        if filtros_desc:
            self.ws['A4'] = f'Filtros: {filtros_desc}'

        return 5

    def adicionar_tabela(self, ranking, linha_inicio=5):
        colunas = ['Fornecedor', 'Total Comprado', '% do Total']

        for col, titulo in enumerate(colunas, 1):
            cell = self.ws.cell(row=linha_inicio, column=col)
            cell.value = titulo
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.border
            cell.alignment = Alignment(horizontal='center', vertical='center')

        linha = linha_inicio + 1
        for item in ranking:
            self.ws.cell(row=linha, column=1).value = item['fornecedor']
            self.ws.cell(row=linha, column=2).value = item['total']
            self.ws.cell(row=linha, column=3).value = item['percentual']

            for col in range(1, 4):
                cell = self.ws.cell(row=linha, column=col)
                cell.border = self.border
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

            linha += 1

        self.ws.column_dimensions['A'].width = 50
        self.ws.column_dimensions['B'].width = 18
        self.ws.column_dimensions['C'].width = 12

        return linha

    def gerar(self):
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
