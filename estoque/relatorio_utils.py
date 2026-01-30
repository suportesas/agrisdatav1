"""
Utilitários para geração de relatórios de estoque em PDF e Excel
"""
import io
from datetime import datetime
from decimal import Decimal
from django.http import HttpResponse
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT, TA_LEFT
from django.db.models import Sum, Q
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from estoque.models import Lote, MovimentacaoEstoque


class RelatorioEstoquePDF:
    """Gerador de relatório de estoque em PDF"""
    
    def __init__(self, empresa, titulo='Relatório de Estoque'):
        self.empresa = empresa
        self.titulo = titulo
        self.buffer = io.BytesIO()
        self.doc = SimpleDocTemplate(
            self.buffer,
            pagesize=A4,
            rightMargin=0.5*inch,
            leftMargin=0.5*inch,
            topMargin=0.75*inch,
            bottomMargin=0.5*inch,
        )
        self.styles = getSampleStyleSheet()
        self.story = []
        self._setup_styles()
    
    def _setup_styles(self):
        """Configurar estilos customizados"""
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=6,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold'
        ))
        
        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#666666'),
            spaceAfter=12,
            alignment=TA_CENTER,
        ))
    
    def adicionar_cabecalho(self, filtros_desc=''):
        """Adicionar cabeçalho ao relatório"""
        self.story.append(Paragraph(self.titulo, self.styles['CustomTitle']))
        
        info_texto = f'<b>Empresa:</b> {self.empresa.nome}<br/>'
        info_texto += f'<b>Gerado em:</b> {datetime.now().strftime("%d/%m/%Y às %H:%M")}<br/>'
        if filtros_desc:
            info_texto += f'<b>Filtros:</b> {filtros_desc}'
        
        self.story.append(Paragraph(info_texto, self.styles['Subtitle']))
        self.story.append(Spacer(1, 0.3*inch))
    
    def adicionar_tabela_lotes(self, lotes):
        """Adicionar tabela de lotes ao relatório"""
        if not lotes.exists():
            self.story.append(Paragraph('Nenhum lote encontrado com os critérios especificados.', self.styles['Normal']))
            return
        
        # Preparar dados
        dados = [['Produto', 'Lote', 'Armazém', 'Qtd. Atual', 'Validade', 'Status']]
        
        for lote in lotes.select_related('produto', 'armazem'):
            validade = lote.data_validade.strftime('%d/%m/%Y') if lote.data_validade else 'S/D'
            dados.append([
                lote.produto.nome[:30],
                lote.numero_lote[:15],
                lote.armazem.nome[:20],
                f'{float(lote.quantidade_atual):.2f} {lote.unidade_medida.sigla}',
                validade,
                lote.get_status_display(),
            ])
        
        # Criar tabela
        tabela = Table(dados, colWidths=[2*inch, 1.2*inch, 1.2*inch, 1.2*inch, 1*inch, 0.8*inch])
        
        # Aplicar estilos
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        
        self.story.append(tabela)

    def adicionar_resumo_produtos(self, lotes):
        """Adicionar resumo de estoque por produto"""
        self.story.append(Spacer(1, 16))
        self.story.append(Paragraph('Resumo por Produto', self.styles['Subtitle']))
        self.story.append(Spacer(1, 8))

        resumo = _gerar_resumo_produtos(lotes, self.empresa)
        if not resumo:
            self.story.append(Paragraph('Nenhum dado para exibir.', self.styles['Normal']))
            return

        dados = [['Produto', 'Est. Atual', 'Entrada', 'Saída']]
        for item in resumo:
            dados.append([
                item['produto'][:40],
                item['estoque_atual'],
                item['entrada'],
                item['saida'],
            ])

        tabela = Table(dados, colWidths=[3.2*inch, 1.2*inch, 1.2*inch, 1.2*inch])
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        self.story.append(tabela)
    
    def gerar(self):
        """Gerar o PDF"""
        self.doc.build(self.story)
        self.buffer.seek(0)
        return self.buffer.getvalue()


class RelatorioEstoqueExcel:
    """Gerador de relatório de estoque em Excel"""
    
    def __init__(self, empresa, titulo='Relatório de Estoque'):
        self.empresa = empresa
        self.titulo = titulo
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Estoque'
        self._setup_styles()
    
    def _setup_styles(self):
        """Configurar estilos do Excel"""
        self.header_fill = PatternFill(start_color='1f4788', end_color='1f4788', fill_type='solid')
        self.header_font = Font(bold=True, color='FFFFFF', size=11)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def adicionar_cabecalho(self, filtros_desc=''):
        """Adicionar cabeçalho"""
        self.ws['A1'] = self.titulo
        self.ws['A1'].font = Font(bold=True, size=14, color='1f4788')
        
        self.ws['A2'] = f'Empresa: {self.empresa.nome}'
        self.ws['A3'] = f'Gerado em: {datetime.now().strftime("%d/%m/%Y às %H:%M")}'
        
        if filtros_desc:
            self.ws['A4'] = f'Filtros: {filtros_desc}'
        
        return 5
    
    def adicionar_tabela_lotes(self, lotes, linha_inicio=5):
        """Adicionar tabela de lotes"""
        # Cabeçalhos
        colunas = ['Produto', 'Lote', 'Armazém', 'Qtd. Atual', 'Unidade', 'Data Fabricação', 'Data Validade', 'Status']
        
        for col, titulo in enumerate(colunas, 1):
            cell = self.ws.cell(row=linha_inicio, column=col)
            cell.value = titulo
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Dados
        linha = linha_inicio + 1
        for lote in lotes.select_related('produto', 'armazem', 'unidade_medida'):
            self.ws.cell(row=linha, column=1).value = lote.produto.nome
            self.ws.cell(row=linha, column=2).value = lote.numero_lote
            self.ws.cell(row=linha, column=3).value = lote.armazem.nome
            self.ws.cell(row=linha, column=4).value = float(lote.quantidade_atual)
            self.ws.cell(row=linha, column=5).value = lote.unidade_medida.sigla
            self.ws.cell(row=linha, column=6).value = lote.data_fabricacao if lote.data_fabricacao else ''
            self.ws.cell(row=linha, column=7).value = lote.data_validade if lote.data_validade else ''
            self.ws.cell(row=linha, column=8).value = lote.get_status_display()
            
            # Aplicar estilos
            for col in range(1, 9):
                cell = self.ws.cell(row=linha, column=col)
                cell.border = self.border
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                
                # Alinhar números à direita
                if col in [4]:
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    cell.number_format = '0.00'
            
            linha += 1
        
        # Ajustar largura das colunas
        self.ws.column_dimensions['A'].width = 30
        self.ws.column_dimensions['B'].width = 15
        self.ws.column_dimensions['C'].width = 20
        self.ws.column_dimensions['D'].width = 12
        self.ws.column_dimensions['E'].width = 10
        self.ws.column_dimensions['F'].width = 15
        self.ws.column_dimensions['G'].width = 15
        self.ws.column_dimensions['H'].width = 15

        return linha

    def adicionar_resumo_produtos(self, lotes, linha_inicio=5):
        """Adicionar resumo por produto em Excel"""
        resumo = _gerar_resumo_produtos(lotes, self.empresa)
        if not resumo:
            self.ws.cell(row=linha_inicio, column=1).value = 'Resumo por Produto'
            self.ws.cell(row=linha_inicio + 1, column=1).value = 'Nenhum dado para exibir.'
            return linha_inicio + 2

        self.ws.cell(row=linha_inicio, column=1).value = 'Resumo por Produto'
        self.ws.cell(row=linha_inicio, column=1).font = Font(bold=True, size=12)

        cabecalhos = ['Produto', 'Est. Atual', 'Entrada', 'Saída']
        linha = linha_inicio + 1
        for col, cab in enumerate(cabecalhos, start=1):
            cell = self.ws.cell(row=linha, column=col)
            cell.value = cab
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.border = self.border
            cell.alignment = Alignment(horizontal='center')

        linha += 1
        for item in resumo:
            self.ws.cell(row=linha, column=1).value = item['produto']
            self.ws.cell(row=linha, column=2).value = item['estoque_atual']
            self.ws.cell(row=linha, column=3).value = item['entrada']
            self.ws.cell(row=linha, column=4).value = item['saida']

            for col in range(1, 5):
                cell = self.ws.cell(row=linha, column=col)
                cell.border = self.border
                cell.alignment = Alignment(horizontal='left')
            linha += 1

        return linha
    
    def gerar(self):
        """Gerar o Excel"""
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


def _gerar_resumo_produtos(lotes, empresa):
    """Gera resumo de estoque por produto com entradas e saídas"""
    if not lotes.exists():
        return []

    produtos_base = list(
        lotes.values('produto_id', 'produto__nome', 'unidade_medida__sigla')
        .annotate(estoque_atual=Sum('quantidade_atual'))
    )

    produto_ids = [p['produto_id'] for p in produtos_base]

    entradas_q = Q(tipo='ENTRADA') | Q(origem__in=['AJUSTE_ENTRADA', 'TRANSFERENCIA_ENTRADA', 'DEVOLUCAO_CLIENTE', 'PRODUCAO', 'COMPRA'])
    saidas_q = Q(tipo='SAIDA') | Q(origem__in=['AJUSTE_SAIDA', 'TRANSFERENCIA_SAIDA', 'DEVOLUCAO_FORNECEDOR', 'REQUISICAO', 'CONSUMO_PRODUCAO', 'VENDA', 'PERDA'])

    entradas = MovimentacaoEstoque.objects.filter(
        empresa=empresa,
        status='EXECUTADO',
        lote__produto_id__in=produto_ids
    ).filter(entradas_q).values('lote__produto_id').annotate(total=Sum('quantidade'))

    saidas = MovimentacaoEstoque.objects.filter(
        empresa=empresa,
        status='EXECUTADO',
        lote__produto_id__in=produto_ids
    ).filter(saidas_q).values('lote__produto_id').annotate(total=Sum('quantidade'))

    entradas_map = {e['lote__produto_id']: e['total'] for e in entradas}
    saidas_map = {s['lote__produto_id']: s['total'] for s in saidas}

    resumo = []
    for item in produtos_base:
        unidade = item['unidade_medida__sigla'] or ''
        estoque = item['estoque_atual'] or 0
        entrada = entradas_map.get(item['produto_id'], 0) or 0
        saida = saidas_map.get(item['produto_id'], 0) or 0

        resumo.append({
            'produto': item['produto__nome'],
            'estoque_atual': f"{float(estoque):.2f} {unidade}".strip(),
            'entrada': f"{float(entrada):.2f} {unidade}".strip(),
            'saida': f"{float(saida):.2f} {unidade}".strip(),
        })

    resumo.sort(key=lambda x: x['produto'])
    return resumo
