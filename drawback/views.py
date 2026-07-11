import io
import json
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import (
    AtoConcessorioForm, DrawbackProdutoACFormSet,
    EntradaDrawbackForm, SaidaDrawbackForm, ImportacaoXLSXForm,
)
from .models import AtoConcessorio, DrawbackProdutoAC, EntradaDrawback, SaidaDrawback


# ── helpers ───────────────────────────────────────────────────────────────────

def _empresa(request):
    return request.user.empresa


@login_required
def ajax_buscar_fornecedores(request):
    """Autocomplete de fornecedores para o campo Despachante Aduaneiro."""
    from django.db.models import Q
    from cadastros.models import Fornecedor
    q = request.GET.get('q', '').strip()
    empresa = _empresa(request)
    if len(q) < 2:
        return JsonResponse({'results': []})
    qs = Fornecedor.objects.filter(empresa=empresa, ativa=True).filter(
        Q(nome__icontains=q) | Q(nome_fantasia__icontains=q) | Q(cnpj_cpf__icontains=q)
    ).order_by('nome')[:15]
    results = []
    for f in qs:
        nome = f.nome_fantasia or f.nome
        cnpj = f.cnpj_cpf
        results.append({
            'id': nome,
            'text': f'{nome} — {cnpj}' if cnpj else nome,
            'nome': nome,
        })
    return JsonResponse({'results': results})


@login_required
def ajax_produto_info(request, produto_id):
    """Retorna NCM e unidade de medida de um produto para auto-preenchimento."""
    from produtos.models import Produto
    empresa = _empresa(request)
    try:
        p = Produto.objects.select_related('ncm', 'unidade_medida').get(
            pk=produto_id, categoria__empresa=empresa, ativo=True
        )
    except Produto.DoesNotExist:
        return JsonResponse({'ok': False}, status=404)
    return JsonResponse({
        'ok': True,
        'ncm': p.ncm.codigo if p.ncm else '',
        'unidade_id': p.unidade_medida_id or '',
        'unidade_sigla': p.unidade_medida.sigla if p.unidade_medida else '',
        'unidade_nome': p.unidade_medida.nome if p.unidade_medida else '',
    })


# ── DASHBOARD ─────────────────────────────────────────────────────────────────

@login_required
def dashboard(request):
    empresa = _empresa(request)
    acs = AtoConcessorio.objects.filter(empresa=empresa).prefetch_related(
        'produtos__entradas', 'produtos__saidas'
    ).order_by('-data_inicio')

    acs_ativos = [ac for ac in acs if ac.status == 'ATIVO']
    acs_vencendo = [ac for ac in acs_ativos if ac.esta_vencendo]

    return render(request, 'drawback/dashboard.html', {
        'acs': acs,
        'acs_ativos': acs_ativos,
        'acs_vencendo': acs_vencendo,
        'sem_acs': len(acs) == 0,
    })


@login_required
def ajuda(request):
    return render(request, 'drawback/ajuda.html')


# ── ATO CONCESSÓRIO ───────────────────────────────────────────────────────────

def _formset_context(empresa, formset):
    """Prepara unidades e linhas de exibição para o formset de produtos."""
    from produtos.models import UnidadeMedida
    unidades = list(UnidadeMedida.objects.filter(empresa=empresa).order_by('sigla'))

    rows = []
    for fform in formset:
        if fform.instance.pk:
            rows.append({
                'prefix': fform.prefix,
                'pk': fform.instance.pk,
                'tipo_produto': fform.instance.tipo_produto,
                'tipo_produto_display': fform.instance.get_tipo_produto_display(),
                'produto_id': fform.instance.produto_id,
                'produto_nome': fform.instance.produto.nome,
                'ncm': fform.instance.ncm or '',
                'descricao': fform.instance.descricao_mercadoria,
                'unidade_id': fform.instance.unidade_medida_id,
                'unidade_sigla': fform.instance.unidade_medida.sigla,
                'qtd_autorizada': str(fform.instance.quantidade_autorizada),
                'qtd_minima': str(fform.instance.quantidade_minima_export or ''),
                'obs': fform.instance.observacoes or '',
            })
        elif fform.has_changed():
            pid = fform['produto'].value() or ''
            produto_nome = ''
            if pid:
                try:
                    p = fform.fields['produto'].queryset.get(pk=pid)
                    produto_nome = p.nome
                except Exception:
                    produto_nome = f'ID {pid}'
            uid = fform['unidade_medida'].value() or ''
            unidade_sigla = ''
            if uid:
                try:
                    u = UnidadeMedida.objects.get(pk=uid)
                    unidade_sigla = u.sigla
                except Exception:
                    pass
            tipo = fform['tipo_produto'].value() or 'MERCADO_INTERNO'
            rows.append({
                'prefix': fform.prefix,
                'pk': '',
                'tipo_produto': tipo,
                'tipo_produto_display': dict(DrawbackProdutoAC.TIPO_PRODUTO_CHOICES).get(tipo, tipo),
                'produto_id': pid,
                'produto_nome': produto_nome,
                'ncm': fform['ncm'].value() or '',
                'descricao': fform['descricao_mercadoria'].value() or '',
                'unidade_id': uid,
                'unidade_sigla': unidade_sigla,
                'qtd_autorizada': fform['quantidade_autorizada'].value() or '',
                'qtd_minima': fform['quantidade_minima_export'].value() or '',
                'obs': fform['observacoes'].value() or '',
            })
    return unidades, rows


@login_required
def ac_create(request):
    empresa = _empresa(request)
    if request.method == 'POST':
        form = AtoConcessorioForm(request.POST)
        formset = DrawbackProdutoACFormSet(request.POST, empresa=empresa)
        if form.is_valid() and formset.is_valid():
            ac = form.save(commit=False)
            ac.empresa = empresa
            ac.criado_por = request.user
            ac.save()
            formset.instance = ac
            formset.save()
            messages.success(request, f'Ato Concessório {ac.numero_ac} cadastrado com sucesso.')
            return redirect('drawback:ac_detail', pk=ac.pk)
    else:
        form = AtoConcessorioForm()
        formset = DrawbackProdutoACFormSet(empresa=empresa)

    unidades, formset_rows = _formset_context(empresa, formset)
    return render(request, 'drawback/form_ac.html', {
        'form': form,
        'formset': formset,
        'titulo': 'Novo Ato Concessório',
        'unidades': unidades,
        'formset_rows': formset_rows,
    })


@login_required
def ac_edit(request, pk):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=pk, empresa=empresa)
    if request.method == 'POST':
        form = AtoConcessorioForm(request.POST, instance=ac)
        formset = DrawbackProdutoACFormSet(request.POST, instance=ac, empresa=empresa)
        if form.is_valid() and formset.is_valid():
            form.save()
            formset.save()
            messages.success(request, 'Ato Concessório atualizado.')
            return redirect('drawback:ac_detail', pk=ac.pk)
    else:
        form = AtoConcessorioForm(instance=ac)
        formset = DrawbackProdutoACFormSet(instance=ac, empresa=empresa)

    unidades, formset_rows = _formset_context(empresa, formset)
    return render(request, 'drawback/form_ac.html', {
        'form': form,
        'formset': formset,
        'ac': ac,
        'titulo': f'Editar AC {ac.numero_ac}',
        'unidades': unidades,
        'formset_rows': formset_rows,
    })


@login_required
def ac_detail(request, pk):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=pk, empresa=empresa)
    produtos = ac.produtos.select_related('produto', 'unidade_medida').order_by('produto__nome')
    entradas = ac.entradas.filter(anulado=False).select_related(
        'drawback_produto__produto', 'fornecedor', 'registrado_por'
    ).order_by('-data_recebimento')
    saidas = ac.saidas.filter(anulado=False).select_related(
        'drawback_produto__produto', 'cliente', 'registrado_por'
    ).order_by('-data_embarque')
    entradas_anuladas = ac.entradas.filter(anulado=True).order_by('-data_registro')
    saidas_anuladas = ac.saidas.filter(anulado=True).order_by('-data_registro')

    resumo_secoes = [
        ('EXPORTACAO',      'Itens de Exportação',            '#16a34a', 'bi bi-box-arrow-up'),
        ('IMPORTACAO',      'Itens de Importação',            '#2563eb', 'bi bi-box-arrow-in-down'),
        ('MERCADO_INTERNO', 'Itens de Compras no Mercado Interno', '#d97706', 'bi bi-shop'),
    ]

    return render(request, 'drawback/detail.html', {
        'ac': ac,
        'produtos': produtos,
        'entradas': entradas,
        'saidas': saidas,
        'entradas_anuladas': entradas_anuladas,
        'saidas_anuladas': saidas_anuladas,
        'resumo_secoes': resumo_secoes,
    })


# ── ENTRADAS ──────────────────────────────────────────────────────────────────

@login_required
def entrada_create(request, ac_pk):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=ac_pk, empresa=empresa)

    if request.method == 'POST':
        form = EntradaDrawbackForm(request.POST, ato_concessorio=ac, empresa=empresa)
        if form.is_valid():
            entrada = form.save(commit=False)
            entrada.ato_concessorio = ac
            entrada.origem = 'MANUAL'
            entrada.retroativo = True
            entrada.registrado_por = request.user
            # Deriva tipo_compra automaticamente do tipo do produto no AC
            if entrada.drawback_produto.tipo_produto == 'IMPORTACAO':
                entrada.tipo_compra = 'IMPORTACAO'
            else:
                entrada.tipo_compra = 'MERCADO_INTERNO'
            entrada.save()
            messages.success(request, 'Entrada registrada com sucesso.')
            return redirect('drawback:ac_detail', pk=ac.pk)
    else:
        form = EntradaDrawbackForm(ato_concessorio=ac, empresa=empresa)

    return render(request, 'drawback/entrada_manual.html', {
        'form': form,
        'ac': ac,
        'titulo': 'Nova Entrada Drawback',
    })


@login_required
@require_POST
def entrada_anular(request, pk):
    empresa = _empresa(request)
    entrada = get_object_or_404(EntradaDrawback, pk=pk, ato_concessorio__empresa=empresa)
    motivo = request.POST.get('motivo_anulacao', '').strip()
    if not motivo:
        messages.error(request, 'Informe o motivo da anulação.')
        return redirect('drawback:ac_detail', pk=entrada.ato_concessorio_id)
    entrada.anulado = True
    entrada.motivo_anulacao = motivo
    entrada.save()
    messages.success(request, 'Entrada anulada.')
    return redirect('drawback:ac_detail', pk=entrada.ato_concessorio_id)


# ── SAÍDAS ────────────────────────────────────────────────────────────────────

@login_required
def saida_create(request, ac_pk):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=ac_pk, empresa=empresa)

    if request.method == 'POST':
        form = SaidaDrawbackForm(request.POST, ato_concessorio=ac, empresa=empresa)
        if form.is_valid():
            saida = form.save(commit=False)
            saida.ato_concessorio = ac
            saida.origem = 'MANUAL'
            saida.retroativo = True
            saida.registrado_por = request.user
            saida.save()
            messages.success(request, 'Saída registrada com sucesso.')
            return redirect('drawback:ac_detail', pk=ac.pk)
    else:
        form = SaidaDrawbackForm(ato_concessorio=ac, empresa=empresa)

    return render(request, 'drawback/saida_manual.html', {
        'form': form,
        'ac': ac,
        'titulo': 'Nova Saída Drawback',
    })


@login_required
@require_POST
def saida_anular(request, pk):
    empresa = _empresa(request)
    saida = get_object_or_404(SaidaDrawback, pk=pk, ato_concessorio__empresa=empresa)
    motivo = request.POST.get('motivo_anulacao', '').strip()
    if not motivo:
        messages.error(request, 'Informe o motivo da anulação.')
        return redirect('drawback:ac_detail', pk=saida.ato_concessorio_id)
    saida.anulado = True
    saida.motivo_anulacao = motivo
    saida.save()
    messages.success(request, 'Saída anulada.')
    return redirect('drawback:ac_detail', pk=saida.ato_concessorio_id)


# ── IMPORTAÇÃO XLSX ───────────────────────────────────────────────────────────

@login_required
def importar_xlsx(request, ac_pk):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=ac_pk, empresa=empresa)

    if request.method == 'POST':
        form = ImportacaoXLSXForm(request.POST, request.FILES)
        if form.is_valid():
            tipo = form.cleaned_data['tipo']
            arquivo = form.cleaned_data['arquivo']
            try:
                import openpyxl
                wb = openpyxl.load_workbook(arquivo, data_only=True)
            except Exception as e:
                messages.error(request, f'Erro ao ler o arquivo: {e}')
                return render(request, 'drawback/importar_xlsx.html', {'form': form, 'ac': ac})

            from .importacao import validar_entradas, validar_saidas
            if tipo == 'entradas':
                resultados, erros_globais = validar_entradas(wb, ac, empresa)
            else:
                resultados, erros_globais = validar_saidas(wb, ac, empresa)

            if erros_globais:
                for e in erros_globais:
                    messages.error(request, e)
                return render(request, 'drawback/importar_xlsx.html', {'form': form, 'ac': ac})

            # Salva resultados na sessão para confirmar
            request.session[f'drawback_import_{ac_pk}'] = {
                'tipo': tipo,
                'resultados': _serialize_resultados(resultados, tipo),
            }

            return render(request, 'drawback/importar_preview.html', {
                'ac': ac,
                'tipo': tipo,
                'resultados': resultados,
                'total_ok': sum(1 for r in resultados if r['ok']),
                'total_erro': sum(1 for r in resultados if not r['ok']),
            })
    else:
        form = ImportacaoXLSXForm()

    return render(request, 'drawback/importar_xlsx.html', {'form': form, 'ac': ac})


def _serialize_resultados(resultados, tipo):
    """Serializa para sessão (apenas campos necessários para gravar)."""
    out = []
    for r in resultados:
        if not r['ok']:
            continue
        if tipo == 'entradas':
            out.append({
                'tipo_compra': r['tipo_compra'],
                'fornecedor_id': r['fornecedor_obj'].pk if r['fornecedor_obj'] else None,
                'fornecedor_str': r['fornecedor_str'],
                'numero_nf': r['numero_nf'],
                'serie_nf': r['serie_nf'],
                'data_nf': r['data_nf'].isoformat() if r['data_nf'] else None,
                'data_recebimento': r['data_recebimento'].isoformat() if r['data_recebimento'] else None,
                'produto_ac_id': r['produto_ac'].pk if r['produto_ac'] else None,
                'produto_variante_id': r['produto_variante'].pk if r.get('produto_variante') else None,
                'quantidade': str(r['quantidade']),
                'valor_unitario': str(r['valor_unitario']) if r['valor_unitario'] else None,
                'observacoes': r['observacoes'],
            })
        else:
            out.append({
                'cliente_id': r['cliente_obj'].pk if r['cliente_obj'] else None,
                'numero_nfe': r['numero_nfe'],
                'numero_re': r['numero_re'],
                'numero_due': r['numero_due'],
                'data_embarque': r['data_embarque'].isoformat() if r['data_embarque'] else None,
                'produto_ac_id': r['produto_ac'].pk if r['produto_ac'] else None,
                'produto_variante_id': r['produto_variante'].pk if r.get('produto_variante') else None,
                'quantidade': str(r['quantidade']),
                'valor_fob_usd': str(r['valor_fob_usd']) if r['valor_fob_usd'] else None,
                'observacoes': r['observacoes'],
            })
    return out


@login_required
@require_POST
def importar_confirmar(request, ac_pk):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=ac_pk, empresa=empresa)

    dados = request.session.pop(f'drawback_import_{ac_pk}', None)
    if not dados:
        messages.error(request, 'Sessão expirada. Faça o upload novamente.')
        return redirect('drawback:importar_xlsx', ac_pk=ac_pk)

    tipo = dados['tipo']
    resultados = dados['resultados']
    criados = 0

    from produtos.models import Produto

    for r in resultados:
        if tipo == 'entradas':
            produto_ac = DrawbackProdutoAC.objects.get(pk=r['produto_ac_id'])
            produto_variante = Produto.objects.get(pk=r['produto_variante_id']) if r.get('produto_variante_id') else None
            from cadastros.models import Fornecedor
            forn = Fornecedor.objects.get(pk=r['fornecedor_id']) if r['fornecedor_id'] else None
            EntradaDrawback.objects.create(
                ato_concessorio=ac,
                drawback_produto=produto_ac,
                produto_variante=produto_variante,
                fornecedor=forn,
                numero_nf=r['numero_nf'],
                serie_nf=r['serie_nf'] or '1',
                data_nf=r['data_nf'],
                data_recebimento=r['data_recebimento'],
                quantidade=Decimal(r['quantidade']),
                valor_unitario=Decimal(r['valor_unitario']) if r['valor_unitario'] else None,
                tipo_compra=r['tipo_compra'],
                origem='MANUAL',
                retroativo=True,
                observacoes=r['observacoes'],
                registrado_por=request.user,
            )
            criados += 1
        else:
            produto_ac = DrawbackProdutoAC.objects.get(pk=r['produto_ac_id'])
            produto_variante = Produto.objects.get(pk=r['produto_variante_id']) if r.get('produto_variante_id') else None
            from cadastros.models import Cliente
            cli = Cliente.objects.get(pk=r['cliente_id']) if r['cliente_id'] else None
            SaidaDrawback.objects.create(
                ato_concessorio=ac,
                drawback_produto=produto_ac,
                produto_variante=produto_variante,
                cliente=cli,
                numero_nfe=r['numero_nfe'],
                numero_re=r['numero_re'],
                numero_due=r['numero_due'],
                data_embarque=r['data_embarque'],
                quantidade=Decimal(r['quantidade']),
                valor_fob_usd=Decimal(r['valor_fob_usd']) if r['valor_fob_usd'] else None,
                origem='MANUAL',
                retroativo=True,
                observacoes=r['observacoes'],
                registrado_por=request.user,
            )
            criados += 1

    tipo_label = 'entradas' if tipo == 'entradas' else 'saídas'
    messages.success(request, f'{criados} {tipo_label} importadas com sucesso.')
    return redirect('drawback:ac_detail', pk=ac.pk)


# ── DOWNLOAD MODELO XLSX ──────────────────────────────────────────────────────

@login_required
def download_modelo_xlsx(request, ac_pk, tipo):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=ac_pk, empresa=empresa)

    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, 'Biblioteca openpyxl não instalada. Execute: pip install openpyxl')
        return redirect('drawback:ac_detail', pk=ac_pk)

    wb = openpyxl.Workbook()

    # Estilos
    roxo = PatternFill(start_color='7C3AED', end_color='7C3AED', fill_type='solid')
    amarelo = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')
    cinza_claro = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
    branco = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
    font_header = Font(bold=True, color='FFFFFF', size=11)
    font_obrig = Font(bold=True, size=10)
    font_exemplo = Font(italic=True, color='9CA3AF', size=10)
    borda = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB'),
    )
    centro = Alignment(horizontal='center', vertical='center')

    produtos_ac = list(ac.produtos.select_related('produto', 'unidade_medida').order_by('produto__nome'))

    if tipo == 'entradas':
        ws = wb.active
        ws.title = 'Entradas'

        cabecalhos = [
            ('Tipo de Compra *', True),
            ('Fornecedor (Nome ou CNPJ) *', True),
            ('Número da NF *', True),
            ('Série da NF', False),
            ('Data da NF * (DD/MM/AAAA)', True),
            ('Data Recebimento * (DD/MM/AAAA)', True),
            ('Produto * (conforme aba Produtos)', True),
            ('Quantidade *', True),
            ('Unidade de Medida *', True),
            ('Valor Unitário R$ (opcional)', False),
            ('Observação (opcional)', False),
        ]

        linha_exemplo = [
            'MERCADO_INTERNO', 'Fornecedor Exemplo LTDA', '000123', '1',
            '15/03/2025', '17/03/2025',
            produtos_ac[0].produto.nome if produtos_ac else 'Nome do Produto',
            '1250.500',
            produtos_ac[0].unidade_medida.sigla if produtos_ac else 'UN',
            '4.50', 'Observação opcional',
        ]

        # Linha 1: título do AC
        ws.merge_cells(f'A1:{get_column_letter(len(cabecalhos))}1')
        ws['A1'] = f'MODELO — ENTRADAS DRAWBACK | AC: {ac.numero_ac} | Vigência: {ac.data_inicio:%d/%m/%Y} a {ac.data_fim_efetiva:%d/%m/%Y}'
        ws['A1'].fill = roxo
        ws['A1'].font = Font(bold=True, color='FFFFFF', size=12)
        ws['A1'].alignment = centro
        ws.row_dimensions[1].height = 28

    else:  # saídas
        ws = wb.active
        ws.title = 'Saídas'

        cabecalhos = [
            ('Cliente (Nome ou CNPJ)', False),
            ('Número NF-e / Invoice', False),
            ('RE (Registro de Exportação)', False),
            ('DUE (Declaração Única de Exportação)', False),
            ('Data Embarque * (DD/MM/AAAA)', True),
            ('Produto * (conforme aba Produtos)', True),
            ('Quantidade *', True),
            ('Unidade de Medida *', True),
            ('Valor FOB USD (opcional)', False),
            ('Observação (opcional)', False),
        ]

        linha_exemplo = [
            'Cliente Exportação LTDA', '000456', '25BR000123456-7', '25BR0012345678901',
            '20/05/2025',
            produtos_ac[0].produto.nome if produtos_ac else 'Nome do Produto',
            '1200.000',
            produtos_ac[0].unidade_medida.sigla if produtos_ac else 'UN',
            '3600.00', 'Embarque Porto de Santos',
        ]

        ws.merge_cells(f'A1:{get_column_letter(len(cabecalhos))}1')
        ws['A1'] = f'MODELO — SAÍDAS DRAWBACK | AC: {ac.numero_ac} | Vigência: {ac.data_inicio:%d/%m/%Y} a {ac.data_fim_efetiva:%d/%m/%Y}'
        ws['A1'].fill = roxo
        ws['A1'].font = Font(bold=True, color='FFFFFF', size=12)
        ws['A1'].alignment = centro
        ws.row_dimensions[1].height = 28

    # Linha 2: cabeçalhos das colunas
    for col, (titulo, obrig) in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=2, column=col, value=titulo)
        cell.fill = roxo
        cell.font = font_header
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = borda
    ws.row_dimensions[2].height = 36

    # Linha 3: exemplo (cinza claro)
    for col, valor in enumerate(linha_exemplo, start=1):
        cell = ws.cell(row=3, column=col, value=valor)
        cell.fill = cinza_claro
        cell.font = font_exemplo
        cell.alignment = Alignment(vertical='center')
        cell.border = borda

    # Linhas 4..53: área de preenchimento com cor por obrigatoriedade
    for row in range(4, 54):
        for col, (_, obrig) in enumerate(cabecalhos, start=1):
            cell = ws.cell(row=row, column=col)
            cell.fill = amarelo if obrig else branco
            cell.border = borda
            cell.alignment = Alignment(vertical='center')

    # Ajusta larguras
    larguras = [22, 32, 16, 8, 22, 22, 36, 12, 18, 18, 30]
    if tipo == 'saidas':
        larguras = [28, 20, 26, 30, 22, 36, 12, 18, 18, 30]
    for col, larg in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col)].width = larg

    ws.freeze_panes = 'A3'

    # ── Aba: Produtos do AC ────────────────────────────────────────────────────
    ws_prod = wb.create_sheet('Produtos do AC')
    ws_prod.merge_cells('A1:E1')
    ws_prod['A1'] = 'PRODUTOS AUTORIZADOS NO AC — use estes nomes exatamente na coluna Produto'
    ws_prod['A1'].fill = roxo
    ws_prod['A1'].font = Font(bold=True, color='FFFFFF', size=11)
    ws_prod['A1'].alignment = centro
    ws_prod.row_dimensions[1].height = 24

    for col, titulo in enumerate(['Nome do Produto', 'NCM', 'Unidade', 'Qtd Autorizada', 'Saldo Atual'], start=1):
        cell = ws_prod.cell(row=2, column=col, value=titulo)
        cell.fill = roxo
        cell.font = font_header
        cell.alignment = centro
        cell.border = borda

    for i, p in enumerate(produtos_ac, start=3):
        ws_prod.cell(row=i, column=1, value=p.produto.nome).border = borda
        ws_prod.cell(row=i, column=2, value=p.ncm).border = borda
        ws_prod.cell(row=i, column=3, value=p.unidade_medida.sigla).border = borda
        ws_prod.cell(row=i, column=4, value=float(p.quantidade_autorizada)).border = borda
        ws_prod.cell(row=i, column=5, value=float(p.saldo_compra if tipo == 'entradas' else p.saldo_exportacao)).border = borda

    for col, larg in enumerate([40, 14, 12, 20, 20], start=1):
        ws_prod.column_dimensions[get_column_letter(col)].width = larg

    # ── Aba: Instruções ────────────────────────────────────────────────────────
    ws_inst = wb.create_sheet('Instruções')
    instrucoes = [
        ('INSTRUÇÕES DE PREENCHIMENTO', True),
        ('', False),
        ('1. Não altere o nome das colunas (linha 2).', False),
        ('2. A linha 3 é um EXEMPLO — substitua pelos seus dados.', False),
        ('3. Campos com fundo AMARELO são obrigatórios.', False),
        ('4. Para o campo "Produto", use exatamente o nome da aba "Produtos do AC".', False),
        ('5. Datas devem estar no formato DD/MM/AAAA.', False),
        ('6. Números decimais: use ponto (.) como separador. Ex: 1250.500', False),
        ('', False),
        ('TIPOS DE COMPRA ACEITOS:', True),
        ('  IMPORTACAO — produto importado do exterior', False),
        ('  MERCADO_INTERNO — produto comprado no Brasil', False),
        ('', False),
        ('DÚVIDAS:', True),
        (f'  AC: {ac.numero_ac} | Vigência: {ac.data_inicio:%d/%m/%Y} a {ac.data_fim_efetiva:%d/%m/%Y}', False),
        ('  Salve o arquivo no formato .xlsx antes de enviar.', False),
    ]
    for i, (texto, negrito) in enumerate(instrucoes, start=1):
        cell = ws_inst.cell(row=i, column=1, value=texto)
        if negrito:
            cell.font = Font(bold=True, color='7C3AED', size=12)
        else:
            cell.font = Font(size=11)
    ws_inst.column_dimensions['A'].width = 70

    # ── Resposta ──────────────────────────────────────────────────────────────
    nome_arquivo = f'modelo_{tipo}_drawback_{ac.numero_ac.replace("/", "-")}.xlsx'
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nome_arquivo}"'
    wb.save(response)
    return response


# ── RELATÓRIO ─────────────────────────────────────────────────────────────────

@login_required
def relatorio(request, ac_pk):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=ac_pk, empresa=empresa)
    produtos = ac.produtos.select_related('produto', 'unidade_medida').order_by('produto__nome')

    entradas_import = ac.entradas.filter(anulado=False, tipo_compra='IMPORTACAO').select_related(
        'drawback_produto__produto', 'fornecedor'
    ).order_by('data_recebimento')

    entradas_mi = ac.entradas.filter(anulado=False, tipo_compra='MERCADO_INTERNO').select_related(
        'drawback_produto__produto', 'fornecedor'
    ).order_by('data_recebimento')

    saidas = ac.saidas.filter(anulado=False).select_related(
        'drawback_produto__produto', 'cliente'
    ).order_by('data_embarque')

    return render(request, 'drawback/relatorio.html', {
        'ac': ac,
        'produtos': produtos,
        'entradas_import': entradas_import,
        'entradas_mi': entradas_mi,
        'saidas': saidas,
    })


@login_required
def relatorio_xlsx(request, ac_pk):
    empresa = _empresa(request)
    ac = get_object_or_404(AtoConcessorio, pk=ac_pk, empresa=empresa)

    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        messages.error(request, 'openpyxl não instalado.')
        return redirect('drawback:relatorio', ac_pk=ac_pk)

    wb = openpyxl.Workbook()
    roxo = PatternFill(start_color='7C3AED', end_color='7C3AED', fill_type='solid')
    azul_claro = PatternFill(start_color='EDE9FE', end_color='EDE9FE', fill_type='solid')
    font_h = Font(bold=True, color='FFFFFF', size=11)
    borda = Border(
        left=Side(style='thin', color='D1D5DB'), right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'), bottom=Side(style='thin', color='D1D5DB'),
    )

    def escrever_secao(ws, titulo, cabecalhos, linhas, row_start):
        ws.cell(row=row_start, column=1, value=titulo).font = Font(bold=True, color='7C3AED', size=12)
        row_start += 1
        for col, h in enumerate(cabecalhos, 1):
            c = ws.cell(row=row_start, column=col, value=h)
            c.fill = roxo; c.font = font_h; c.border = borda
        row_start += 1
        for linha in linhas:
            for col, val in enumerate(linha, 1):
                c = ws.cell(row=row_start, column=col, value=val)
                c.border = borda
            row_start += 1
        return row_start + 1

    # Aba Resumo
    ws = wb.active
    ws.title = 'Resumo'
    ws['A1'] = f'RELATÓRIO DE COMPROVAÇÃO — AC {ac.numero_ac}'
    ws['A1'].font = Font(bold=True, size=14, color='7C3AED')
    ws['A2'] = f'Vigência: {ac.data_inicio:%d/%m/%Y} a {ac.data_fim_efetiva:%d/%m/%Y} | Tipo: {ac.get_tipo_display()}'
    ws['A2'].font = Font(size=11)

    produtos = list(ac.produtos.select_related('produto', 'unidade_medida').order_by('produto__nome'))
    linhas_resumo = [
        [p.produto.nome, p.unidade_medida.sigla,
         float(p.quantidade_autorizada), float(p.quantidade_comprada),
         float(p.saldo_compra), float(p.quantidade_exportada),
         float(p.saldo_exportacao), float(p.percentual_cumprido)]
        for p in produtos
    ]
    escrever_secao(ws, 'RESUMO POR PRODUTO', [
        'Produto', 'UN', 'Autorizado', 'Comprado', 'Saldo Compra',
        'Exportado', 'Saldo Export', '% Cumprido',
    ], linhas_resumo, row_start=4)
    ws.column_dimensions['A'].width = 36

    # Aba Entradas
    ws_e = wb.create_sheet('Entradas')
    entradas = list(ac.entradas.filter(anulado=False).select_related(
        'drawback_produto__produto', 'fornecedor').order_by('data_recebimento'))
    linhas_e = [
        [e.get_tipo_compra_display(), str(e.fornecedor) if e.fornecedor else '',
         e.numero_nf, e.data_nf.strftime('%d/%m/%Y') if e.data_nf else '',
         e.data_recebimento.strftime('%d/%m/%Y'),
         e.drawback_produto.produto.nome, float(e.quantidade),
         float(e.valor_unitario) if e.valor_unitario else '',
         '✓' if e.retroativo else '']
        for e in entradas
    ]
    escrever_secao(ws_e, 'ENTRADAS DE INSUMOS', [
        'Tipo', 'Fornecedor', 'NF', 'Data NF', 'Data Recebimento',
        'Produto', 'Quantidade', 'Valor Unit R$', 'Retroativo',
    ], linhas_e, row_start=1)

    # Aba Saídas
    ws_s = wb.create_sheet('Saídas')
    saidas = list(ac.saidas.filter(anulado=False).select_related(
        'drawback_produto__produto', 'cliente').order_by('data_embarque'))
    linhas_s = [
        [str(s.cliente) if s.cliente else '', s.numero_nfe, s.numero_re, s.numero_due,
         s.data_embarque.strftime('%d/%m/%Y'),
         s.drawback_produto.produto.nome, float(s.quantidade),
         float(s.valor_fob_usd) if s.valor_fob_usd else '',
         '✓' if s.retroativo else '']
        for s in saidas
    ]
    escrever_secao(ws_s, 'EXPORTAÇÕES REALIZADAS', [
        'Cliente', 'NF-e', 'RE', 'DUE', 'Data Embarque',
        'Produto', 'Quantidade', 'FOB USD', 'Retroativo',
    ], linhas_s, row_start=1)

    nome = f'relatorio_drawback_{ac.numero_ac.replace("/", "-")}.xlsx'
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nome}"'
    wb.save(response)
    return response


# ── VÍNCULOS VIA ERP ──────────────────────────────────────────────────────────

@login_required
def vincular_entrada_erp(request, item_recebimento_id):
    from compras.models import ItemRecebimento
    from . import services

    empresa = _empresa(request)
    item = get_object_or_404(ItemRecebimento, pk=item_recebimento_id,
                             recebimento__empresa=empresa)

    if hasattr(item, 'drawback_entrada'):
        messages.warning(request, 'Este item já está vinculado ao Drawback.')
        return redirect('compras:recebimento_detail', pk=item.recebimento_id)

    produto = services.produto_do_recebimento(item)
    sugestao = services.sugerir_vinculo_entrada(item, empresa) if produto else None

    ac_id = request.POST.get('ac_id') or request.GET.get('ac_id')
    if ac_id:
        ac = get_object_or_404(AtoConcessorio, pk=ac_id, empresa=empresa, status='ATIVO')
    elif sugestao:
        ac = sugestao.ato_concessorio
    else:
        ac = None

    acs_candidatos = AtoConcessorio.objects.filter(
        empresa=empresa, status='ATIVO', produtos__produto=produto,
        produtos__tipo_produto__in=['IMPORTACAO', 'MERCADO_INTERNO'],
    ).distinct() if produto else AtoConcessorio.objects.none()

    if ac is None:
        # Sem sugestão automática: usuário escolhe o AC manualmente.
        return render(request, 'drawback/vincular_entrada_erp.html', {
            'item': item,
            'acs': AtoConcessorio.objects.filter(empresa=empresa, status='ATIVO'),
            'acs_candidatos': acs_candidatos,
            'sem_sugestao': True,
        })

    fornecedor = item.recebimento.fornecedor or (
        item.recebimento.pedido.fornecedor if item.recebimento.pedido_id else None
    )
    # Se o produto do recebimento é uma variante, ela já é conhecida com certeza —
    # não depende de escolha manual do usuário.
    variante_recebimento = produto if produto and produto.produto_base_id else None

    initial = {
        'drawback_produto': sugestao.pk if sugestao and sugestao.ato_concessorio_id == ac.pk else None,
        'produto_variante': variante_recebimento.pk if variante_recebimento else None,
        'fornecedor': fornecedor.pk if fornecedor else None,
        'numero_nf': item.recebimento.nota_fiscal,
        'serie_nf': item.recebimento.serie_nota,
        'data_nf': item.recebimento.data_emissao_nf,
        'data_recebimento': item.recebimento.data_recebimento,
        'quantidade': min(item.quantidade_recebida, sugestao.saldo_compra) if sugestao else item.quantidade_recebida,
    }

    if request.method == 'POST':
        form = EntradaDrawbackForm(request.POST, ato_concessorio=ac, empresa=empresa)
        if form.is_valid():
            entrada = form.save(commit=False)
            entrada.ato_concessorio = ac
            entrada.item_recebimento = item
            entrada.produto_variante = variante_recebimento
            entrada.origem = 'ERP'
            entrada.retroativo = False
            entrada.registrado_por = request.user
            entrada.tipo_compra = (
                'IMPORTACAO' if entrada.drawback_produto.tipo_produto == 'IMPORTACAO' else 'MERCADO_INTERNO'
            )
            entrada.save()
            messages.success(request, 'Item vinculado ao Drawback com sucesso.')
            return redirect('compras:recebimento_detail', pk=item.recebimento_id)
    else:
        form = EntradaDrawbackForm(ato_concessorio=ac, empresa=empresa, initial=initial)

    return render(request, 'drawback/vincular_entrada_erp.html', {
        'form': form,
        'item': item,
        'ac': ac,
        'sugestao': sugestao,
        'acs_candidatos': acs_candidatos,
    })


@login_required
def vincular_saida_erp(request, item_nfe_id):
    from faturamento.models import ItemNotaFiscal
    from . import services

    empresa = _empresa(request)
    item_nfe = get_object_or_404(ItemNotaFiscal, pk=item_nfe_id, nota_fiscal__empresa=empresa)

    if hasattr(item_nfe, 'drawback_saida'):
        messages.warning(request, 'Este item já está vinculado ao Drawback.')
        return redirect('faturamento:nota_fiscal_detail', pk=item_nfe.nota_fiscal_id)

    sugestao = services.sugerir_vinculo_saida(item_nfe, empresa)

    ac_id = request.POST.get('ac_id') or request.GET.get('ac_id')
    if ac_id:
        ac = get_object_or_404(AtoConcessorio, pk=ac_id, empresa=empresa, status='ATIVO')
    elif sugestao:
        ac = sugestao.ato_concessorio
    else:
        ac = None

    acs_candidatos = AtoConcessorio.objects.filter(
        empresa=empresa, status='ATIVO', produtos__produto_id=item_nfe.produto_id,
        produtos__tipo_produto='EXPORTACAO',
    ).distinct() if item_nfe.produto_id else AtoConcessorio.objects.none()

    if ac is None:
        return render(request, 'drawback/vincular_saida_erp.html', {
            'item_nfe': item_nfe,
            'acs': AtoConcessorio.objects.filter(empresa=empresa, status='ATIVO'),
            'acs_candidatos': acs_candidatos,
            'sem_sugestao': True,
        })

    variante_nfe = item_nfe.produto if item_nfe.produto_id and item_nfe.produto.produto_base_id else None

    initial = {
        'drawback_produto': sugestao.pk if sugestao and sugestao.ato_concessorio_id == ac.pk else None,
        'produto_variante': variante_nfe.pk if variante_nfe else None,
        'cliente': item_nfe.nota_fiscal.cliente_id,
        'numero_nfe': item_nfe.nota_fiscal.numero,
        'data_embarque': timezone.now().date(),
        'quantidade': min(item_nfe.quantidade, sugestao.saldo_exportacao) if sugestao else item_nfe.quantidade,
    }

    if request.method == 'POST':
        form = SaidaDrawbackForm(request.POST, ato_concessorio=ac, empresa=empresa)
        if form.is_valid():
            saida = form.save(commit=False)
            saida.ato_concessorio = ac
            saida.item_nfe = item_nfe
            saida.produto_variante = variante_nfe
            saida.origem = 'ERP'
            saida.retroativo = False
            saida.registrado_por = request.user
            if not saida.numero_nfe:
                saida.numero_nfe = item_nfe.nota_fiscal.numero
            saida.save()
            messages.success(request, 'NF-e vinculada ao Drawback com sucesso.')
            return redirect('faturamento:nota_fiscal_detail', pk=item_nfe.nota_fiscal_id)
    else:
        form = SaidaDrawbackForm(ato_concessorio=ac, empresa=empresa, initial=initial)

    return render(request, 'drawback/vincular_saida_erp.html', {
        'form': form,
        'item_nfe': item_nfe,
        'ac': ac,
        'sugestao': sugestao,
        'acs_candidatos': acs_candidatos,
    })


# ── AJAX ──────────────────────────────────────────────────────────────────────

@login_required
def ajax_saldo_produto(request, produto_ac_id):
    empresa = _empresa(request)
    produto_ac = get_object_or_404(
        DrawbackProdutoAC,
        pk=produto_ac_id,
        ato_concessorio__empresa=empresa,
    )
    variantes = [
        {'id': v.pk, 'nome_completo': v.nome_completo}
        for v in produto_ac.produto.variantes.filter(ativo=True).order_by('nome')
    ]
    return JsonResponse({
        'produto': produto_ac.produto.nome,
        'tipo_produto': produto_ac.tipo_produto,
        'unidade': produto_ac.unidade_medida.sigla,
        'quantidade_autorizada': float(produto_ac.quantidade_autorizada),
        'quantidade_comprada': float(produto_ac.quantidade_comprada),
        'saldo_compra': float(produto_ac.saldo_compra),
        'quantidade_exportada': float(produto_ac.quantidade_exportada),
        'saldo_exportacao': float(produto_ac.saldo_exportacao),
        'percentual_cumprido': float(produto_ac.percentual_cumprido),
        'variantes': variantes,
    })
