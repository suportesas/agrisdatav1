"""
Utilitários para geração de relatório de compras por período
"""
import io
from datetime import datetime
from decimal import Decimal
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class RelatorioComprasPDF:
    """Gerador de relatório de compras em PDF"""
    
    def __init__(self, empresa, titulo='Relatório de Compras'):
        self.empresa = empresa
        self.titulo = titulo
        self.buffer = io.BytesIO()
        self.doc = SimpleDocTemplate(
            self.buffer,
            pagesize=landscape(A4),  # Paisagem para mais colunas
            rightMargin=0.5*inch,
            leftMargin=0.5*inch,
            topMargin=0.75*inch,
            bottomMargin=0.5*inch,
        )
        self.styles = getSampleStyleSheet()
        self.story = []
        self._setup_styles()
    
    def _setup_styles(self):
        """Configurar estilos customizados"""
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=6,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold'
        ))
        
        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#666666'),
            spaceAfter=12,
            alignment=TA_CENTER,
        ))
    
    def adicionar_cabecalho(self, filtros_desc=''):
        """Adicionar cabeçalho ao relatório"""
        self.story.append(Paragraph(self.titulo, self.styles['CustomTitle']))
        
        info_texto = f'<b>Empresa:</b> {self.empresa.nome}<br/>'
        info_texto += f'<b>Gerado em:</b> {datetime.now().strftime("%d/%m/%Y às %H:%M")}<br/>'
        if filtros_desc:
            info_texto += f'<b>Filtros:</b> {filtros_desc}'
        
        self.story.append(Paragraph(info_texto, self.styles['Subtitle']))
        self.story.append(Spacer(1, 0.3*inch))
    
    def adicionar_tabela_compras(self, compras_data, agrupar_por_fornecedor=False):
        """Adicionar tabela de compras ao relatório"""
        if not compras_data:
            self.story.append(Paragraph('Nenhuma compra encontrada com os critérios especificados.', self.styles['Normal']))
            return
        
        if agrupar_por_fornecedor:
            self._adicionar_tabela_agrupada(compras_data)
        else:
            self._adicionar_tabela_simples(compras_data)
    
    def _adicionar_tabela_simples(self, compras_data):
        """Tabela simples sem agrupamento"""
        dados = [['Pedido', 'Nota Fiscal', 'Data Emissão', 'Fornecedor', 'Produto', 'Qtd', 'Valor Unit.', 'Valor Total']]
        
        for item in compras_data:
            dados.append([
                item['pedido'][:12],
                item['nota_fiscal'][:15],
                item['data_emissao'],
                item['fornecedor'][:20],
                item['produto'][:25],
                item['quantidade'],
                item['valor_unitario'],
                item['valor_total'],
            ])
        
        # Criar tabela
        tabela = Table(dados, colWidths=[0.9*inch, 1*inch, 0.9*inch, 1.5*inch, 2*inch, 0.7*inch, 0.9*inch, 1*inch])
        
        # Aplicar estilos
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('LEFTPADDING', (0, 0), (-1, -1), 3),
            ('RIGHTPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        
        self.story.append(tabela)
    
    def _adicionar_tabela_agrupada(self, compras_data):
        """Tabela agrupada por fornecedor"""
        # Agrupar por fornecedor
        por_fornecedor = {}
        for item in compras_data:
            fornecedor = item['fornecedor']
            if fornecedor not in por_fornecedor:
                por_fornecedor[fornecedor] = []
            por_fornecedor[fornecedor].append(item)
        
        for fornecedor, itens in sorted(por_fornecedor.items()):
            # Adicionar título do fornecedor
            self.story.append(Paragraph(f'<b>Fornecedor: {fornecedor}</b>', self.styles['Heading3']))
            self.story.append(Spacer(1, 0.1*inch))
            
            # Tabela de itens
            dados = [['Pedido', 'Nota Fiscal', 'Data Emissão', 'Produto', 'Qtd', 'Valor Unit.', 'Valor Total']]
            
            total_fornecedor = 0
            for item in itens:
                dados.append([
                    item['pedido'][:12],
                    item['nota_fiscal'][:15],
                    item['data_emissao'],
                    item['produto'][:30],
                    item['quantidade'],
                    item['valor_unitario'],
                    item['valor_total'],
                ])
                total_fornecedor += item['valor_total_num']
            
            # Linha de total
            dados.append(['', '', '', 'TOTAL', '', '', f'R$ {total_fornecedor:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')])
            
            tabela = Table(dados, colWidths=[0.9*inch, 1*inch, 0.9*inch, 2.5*inch, 0.7*inch, 0.9*inch, 1*inch])
            
            tabela.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#5a6268')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 8),
                ('GRID', (0, 0), (-1, -2), 1, colors.black),
                ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f5f5f5')]),
                ('FONTSIZE', (0, 1), (-1, -1), 7),
                # Total
                ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e9ecef')),
                ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
                ('SPAN', (0, -1), (2, -1)),
                ('LEFTPADDING', (0, 0), (-1, -1), 3),
                ('RIGHTPADDING', (0, 0), (-1, -1), 3),
            ]))
            
            self.story.append(tabela)
            self.story.append(Spacer(1, 0.2*inch))
    
    def gerar(self):
        """Gerar o PDF"""
        self.doc.build(self.story)
        self.buffer.seek(0)
        return self.buffer.getvalue()


class RelatorioComprasExcel:
    """Gerador de relatório de compras em Excel"""
    
    def __init__(self, empresa, titulo='Relatório de Compras'):
        self.empresa = empresa
        self.titulo = titulo
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Compras'
        self._setup_styles()
    
    def _setup_styles(self):
        """Configurar estilos do Excel"""
        self.header_fill = PatternFill(start_color='1f4788', end_color='1f4788', fill_type='solid')
        self.header_font = Font(bold=True, color='FFFFFF', size=11)
        self.group_fill = PatternFill(start_color='5a6268', end_color='5a6268', fill_type='solid')
        self.group_font = Font(bold=True, color='FFFFFF', size=11)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def adicionar_cabecalho(self, filtros_desc=''):
        """Adicionar cabeçalho"""
        self.ws['A1'] = self.titulo
        self.ws['A1'].font = Font(bold=True, size=14, color='1f4788')
        
        self.ws['A2'] = f'Empresa: {self.empresa.nome}'
        self.ws['A3'] = f'Gerado em: {datetime.now().strftime("%d/%m/%Y às %H:%M")}'
        
        if filtros_desc:
            self.ws['A4'] = f'Filtros: {filtros_desc}'
        
        return 5
    
    def adicionar_tabela_compras(self, compras_data, linha_inicio=5, agrupar_por_fornecedor=False):
        """Adicionar tabela de compras"""
        if agrupar_por_fornecedor:
            return self._adicionar_tabela_agrupada(compras_data, linha_inicio)
        else:
            return self._adicionar_tabela_simples(compras_data, linha_inicio)
    
    def _adicionar_tabela_simples(self, compras_data, linha_inicio):
        """Tabela simples"""
        # Cabeçalhos
        colunas = ['Pedido', 'Nota Fiscal', 'Data Emissão', 'Fornecedor', 'Produto', 'Quantidade', 'Valor Unitário', 'Valor Total']
        
        for col, titulo in enumerate(colunas, 1):
            cell = self.ws.cell(row=linha_inicio, column=col)
            cell.value = titulo
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Dados
        linha = linha_inicio + 1
        for item in compras_data:
            self.ws.cell(row=linha, column=1).value = item['pedido']
            self.ws.cell(row=linha, column=2).value = item['nota_fiscal']
            self.ws.cell(row=linha, column=3).value = item['data_emissao']
            self.ws.cell(row=linha, column=4).value = item['fornecedor']
            self.ws.cell(row=linha, column=5).value = item['produto']
            self.ws.cell(row=linha, column=6).value = item['quantidade_num']
            self.ws.cell(row=linha, column=7).value = item['valor_unitario_num']
            self.ws.cell(row=linha, column=8).value = item['valor_total_num']
            
            # Aplicar estilos
            for col in range(1, 9):
                cell = self.ws.cell(row=linha, column=col)
                cell.border = self.border
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
                if col in [6]:
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    cell.number_format = '0.00'
                elif col in [7, 8]:
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    cell.number_format = 'R$ #,##0.00'
            
            linha += 1
        
        # Ajustar largura das colunas
        self.ws.column_dimensions['A'].width = 15
        self.ws.column_dimensions['B'].width = 18
        self.ws.column_dimensions['C'].width = 15
        self.ws.column_dimensions['D'].width = 30
        self.ws.column_dimensions['E'].width = 35
        self.ws.column_dimensions['F'].width = 12
        self.ws.column_dimensions['G'].width = 15
        self.ws.column_dimensions['H'].width = 15
        
        return linha
    
    def _adicionar_tabela_agrupada(self, compras_data, linha_inicio):
        """Tabela agrupada por fornecedor"""
        # Agrupar por fornecedor
        por_fornecedor = {}
        for item in compras_data:
            fornecedor = item['fornecedor']
            if fornecedor not in por_fornecedor:
                por_fornecedor[fornecedor] = []
            por_fornecedor[fornecedor].append(item)
        
        linha = linha_inicio
        
        for fornecedor, itens in sorted(por_fornecedor.items()):
            # Cabeçalho do fornecedor
            self.ws.cell(row=linha, column=1).value = f'Fornecedor: {fornecedor}'
            self.ws.cell(row=linha, column=1).font = self.group_font
            self.ws.cell(row=linha, column=1).fill = self.group_fill
            self.ws.merge_cells(f'A{linha}:H{linha}')
            linha += 1
            
            # Cabeçalhos das colunas
            colunas = ['Pedido', 'Nota Fiscal', 'Data Emissão', 'Produto', 'Quantidade', 'Valor Unitário', 'Valor Total']
            for col, titulo in enumerate(colunas, 1):
                cell = self.ws.cell(row=linha, column=col)
                cell.value = titulo
                cell.fill = self.header_fill
                cell.font = self.header_font
                cell.border = self.border
                cell.alignment = Alignment(horizontal='center', vertical='center')
            linha += 1
            
            # Dados do fornecedor
            total_fornecedor = 0
            for item in itens:
                self.ws.cell(row=linha, column=1).value = item['pedido']
                self.ws.cell(row=linha, column=2).value = item['nota_fiscal']
                self.ws.cell(row=linha, column=3).value = item['data_emissao']
                self.ws.cell(row=linha, column=4).value = item['produto']
                self.ws.cell(row=linha, column=5).value = item['quantidade_num']
                self.ws.cell(row=linha, column=6).value = item['valor_unitario_num']
                self.ws.cell(row=linha, column=7).value = item['valor_total_num']
                
                for col in range(1, 8):
                    cell = self.ws.cell(row=linha, column=col)
                    cell.border = self.border
                    if col in [5]:
                        cell.number_format = '0.00'
                        cell.alignment = Alignment(horizontal='right')
                    elif col in [6, 7]:
                        cell.number_format = 'R$ #,##0.00'
                        cell.alignment = Alignment(horizontal='right')
                
                total_fornecedor += item['valor_total_num']
                linha += 1
            
            # Total do fornecedor
            self.ws.cell(row=linha, column=4).value = 'TOTAL'
            self.ws.cell(row=linha, column=4).font = Font(bold=True)
            self.ws.cell(row=linha, column=7).value = total_fornecedor
            self.ws.cell(row=linha, column=7).font = Font(bold=True)
            self.ws.cell(row=linha, column=7).number_format = 'R$ #,##0.00'
            self.ws.cell(row=linha, column=7).alignment = Alignment(horizontal='right')
            linha += 2  # Espaço entre fornecedores
        
        # Ajustar largura
        self.ws.column_dimensions['A'].width = 15
        self.ws.column_dimensions['B'].width = 18
        self.ws.column_dimensions['C'].width = 15
        self.ws.column_dimensions['D'].width = 35
        self.ws.column_dimensions['E'].width = 12
        self.ws.column_dimensions['F'].width = 15
        self.ws.column_dimensions['G'].width = 15
        
        return linha
    
    def gerar(self):
        """Gerar o Excel"""
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


def gerar_dados_compras(pedidos, ordenacao='fornecedor'):
    """
    Gera dados de compras para o relatório
    Extrai informações dos pedidos e seus recebimentos
    
    Args:
        pedidos: QuerySet de PedidoCompra
        ordenacao: Campo para ordenação ('fornecedor', 'nota_fiscal', 'data_emissao' ou 'categoria')
    """
    dados = []
    
    for pedido in pedidos:
        # Buscar recebimentos do pedido
        recebimentos = pedido.recebimentos.filter(status='RECEBIDO').prefetch_related('itens__produto')
        
        if recebimentos.exists():
            # Se tem recebimentos, usa dados dos recebimentos (com nota fiscal)
            for recebimento in recebimentos:
                for item_receb in recebimento.itens.all():
                    dados.append({
                        'pedido': pedido.numero,
                        'nota_fiscal': recebimento.nota_fiscal or 'S/NF',
                        'data_emissao': recebimento.data_recebimento.strftime('%d/%m/%Y'),
                        'data_emissao_sort': recebimento.data_recebimento,
                        'fornecedor': pedido.fornecedor.nome,
                        'produto': item_receb.produto.nome,
                        'categoria': item_receb.produto.categoria.nome if item_receb.produto.categoria else 'Sem Categoria',
                        'quantidade': f'{float(item_receb.quantidade_recebida):.2f}',
                        'quantidade_num': float(item_receb.quantidade_recebida),
                        'valor_unitario': f'R$ {float(item_receb.valor_unitario):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.'),
                        'valor_unitario_num': float(item_receb.valor_unitario),
                        'valor_total': f'R$ {float(item_receb.valor_total):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.'),
                        'valor_total_num': float(item_receb.valor_total),
                    })
        else:
            # Se não tem recebimento, usa dados do pedido
            for item_pedido in pedido.itens.all():
                dados.append({
                    'pedido': pedido.numero,
                    'nota_fiscal': 'Pendente',
                    'data_emissao': pedido.data_pedido.strftime('%d/%m/%Y'),
                    'data_emissao_sort': pedido.data_pedido,
                    'fornecedor': pedido.fornecedor.nome,
                    'produto': item_pedido.produto.nome,
                    'categoria': item_pedido.produto.categoria.nome if item_pedido.produto.categoria else 'Sem Categoria',
                    'quantidade': f'{float(item_pedido.quantidade_pedida):.2f}',
                    'quantidade_num': float(item_pedido.quantidade_pedida),
                    'valor_unitario': f'R$ {float(item_pedido.preco_unitario):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.'),
                    'valor_unitario_num': float(item_pedido.preco_unitario),
                    'valor_total': f'R$ {float(item_pedido.preco_total):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.'),
                    'valor_total_num': float(item_pedido.preco_total),
                })
    
    # Ordenar dados conforme selecionado
    if ordenacao == 'fornecedor':
        dados.sort(key=lambda x: x['fornecedor'])
    elif ordenacao == 'nota_fiscal':
        dados.sort(key=lambda x: x['nota_fiscal'])
    elif ordenacao == 'data_emissao':
        dados.sort(key=lambda x: x['data_emissao_sort'])
    elif ordenacao == 'categoria':
        dados.sort(key=lambda x: x['categoria'])
    
    return dados
