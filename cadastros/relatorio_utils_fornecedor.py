"""
Utilitarios para geracao de relatorios de fornecedores em PDF e Excel
"""
import io
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.utils import ImageReader
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def _create_logo_cell(logo_field, max_width, max_height):
    try:
        try:
            logo_source = logo_field.path
        except NotImplementedError:
            # Storage remoto (ex: Supabase) nao suporta path local: baixa via URL
            import urllib.request
            logo_source = io.BytesIO(urllib.request.urlopen(logo_field.url, timeout=5).read())
        reader = ImageReader(logo_source)
        img_width, img_height = reader.getSize()
        if not img_width or not img_height:
            return ''
        scale = min(max_width / img_width, max_height / img_height)
        if hasattr(logo_source, 'seek'):
            logo_source.seek(0)
        return Image(logo_source, width=img_width * scale, height=img_height * scale)
    except Exception:
        return ''


class RelatorioFornecedoresPDF:
    def __init__(self, empresa, titulo='Relatorio de Fornecedores'):
        self.empresa = empresa
        self.titulo = titulo
        self.buffer = io.BytesIO()
        self.doc = SimpleDocTemplate(
            self.buffer,
            pagesize=A4,
            rightMargin=0.5 * inch,
            leftMargin=0.5 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.5 * inch,
        )
        self.styles = getSampleStyleSheet()
        self.story = []
        self._setup_styles()

    def _setup_styles(self):
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=6,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold'
        ))

        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#666666'),
            spaceAfter=12,
            alignment=TA_CENTER,
        ))

    def adicionar_cabecalho(self, filtros_desc=''):
        info_texto = f'<b>Empresa:</b> {self.empresa.nome}<br/>'
        info_texto += f'<b>Gerado em:</b> {datetime.now().strftime("%d/%m/%Y as %H:%M")}<br/>'
        if filtros_desc:
            info_texto += f'<b>Filtros:</b> {filtros_desc}'

        logo_cell = ''
        try:
            if getattr(self.empresa, 'logo', None) and self.empresa.logo.name:
                logo_cell = _create_logo_cell(self.empresa.logo, 1.6 * inch, 1.0 * inch)
        except Exception:
            logo_cell = ''

        header_table = Table(
            [[logo_cell, Paragraph(info_texto, self.styles['Subtitle'])]],
            colWidths=[1.8 * inch, 4.8 * inch]
        )
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))

        self.story.append(header_table)
        self.story.append(Spacer(1, 0.12 * inch))
        self.story.append(Paragraph(self.titulo, self.styles['CustomTitle']))
        self.story.append(Spacer(1, 0.3 * inch))

    def adicionar_tabela(self, fornecedores):
        if not fornecedores.exists():
            self.story.append(Paragraph('Nenhum fornecedor encontrado com os criterios especificados.', self.styles['Normal']))
            return

        dados = [['Razao Social', 'CNPJ/CPF', 'Cidade/UF', 'Status']]

        for fornecedor in fornecedores:
            cnpj_cpf = fornecedor.get_cnpj_cpf_formatado() if hasattr(fornecedor, 'get_cnpj_cpf_formatado') else (fornecedor.cnpj_cpf or '')
            cidade_uf = ''
            if fornecedor.cidade and fornecedor.uf:
                cidade_uf = f"{fornecedor.cidade}/{fornecedor.uf}"
            elif fornecedor.cidade:
                cidade_uf = fornecedor.cidade
            elif fornecedor.uf:
                cidade_uf = fornecedor.uf
            status = 'Ativo' if fornecedor.ativa else 'Inativo'
            dados.append([
                (fornecedor.nome or '')[:50],
                cnpj_cpf,
                cidade_uf,
                status,
            ])

        tabela = Table(dados, colWidths=[3.0 * inch, 1.5 * inch, 1.6 * inch, 0.9 * inch])
        tabela.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f4788')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))

        self.story.append(tabela)

    def gerar(self):
        self.doc.build(self.story)
        self.buffer.seek(0)
        return self.buffer.getvalue()


class RelatorioFornecedorPDF:
    def __init__(self, empresa, fornecedor, titulo='Cadastro de Fornecedor'):
        self.empresa = empresa
        self.fornecedor = fornecedor
        self.titulo = titulo
        self.buffer = io.BytesIO()
        self.doc = SimpleDocTemplate(
            self.buffer,
            pagesize=A4,
            rightMargin=0.5 * inch,
            leftMargin=0.5 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.5 * inch,
        )
        self.styles = getSampleStyleSheet()
        self.story = []
        self._setup_styles()

    def _setup_styles(self):
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1f4788'),
            spaceAfter=6,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold'
        ))

        self.styles.add(ParagraphStyle(
            name='Subtitle',
            parent=self.styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#666666'),
            spaceAfter=12,
            alignment=TA_CENTER,
        ))

        self.styles.add(ParagraphStyle(
            name='SectionTitle',
            parent=self.styles['Heading2'],
            fontSize=12,
            textColor=colors.HexColor('#1f4788'),
            spaceBefore=10,
            spaceAfter=6,
            fontName='Helvetica-Bold'
        ))

    def _fmt(self, value, default='—'):
        if value is None or value == '':
            return default
        return str(value)

    def adicionar_cabecalho(self):
        info_texto = f'<b>Empresa:</b> {self.empresa.nome}<br/>'
        info_texto += f'<b>Gerado em:</b> {datetime.now().strftime("%d/%m/%Y as %H:%M")}<br/>'
        info_texto += f'<b>Fornecedor:</b> {self.fornecedor.get_nome_exibicao()}'

        logo_cell = ''
        try:
            if getattr(self.empresa, 'logo', None) and self.empresa.logo.name:
                logo_cell = _create_logo_cell(self.empresa.logo, 1.6 * inch, 1.0 * inch)
        except Exception:
            logo_cell = ''

        header_table = Table(
            [[logo_cell, Paragraph(info_texto, self.styles['Subtitle'])]],
            colWidths=[1.8 * inch, 4.8 * inch]
        )
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))

        self.story.append(header_table)
        self.story.append(Spacer(1, 0.12 * inch))
        self.story.append(Paragraph(self.titulo, self.styles['CustomTitle']))
        self.story.append(Spacer(1, 0.2 * inch))

    def _add_kv_table(self, titulo, linhas):
        self.story.append(Paragraph(titulo, self.styles['SectionTitle']))

        dados = []
        for chave, valor in linhas:
            dados.append([
                Paragraph(f'<b>{chave}</b>', self.styles['Normal']),
                Paragraph(self._fmt(valor), self.styles['Normal'])
            ])

        tabela = Table(dados, colWidths=[2.2 * inch, 4.8 * inch])
        tabela.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#c7c7c7')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))

        self.story.append(tabela)

    def _add_list(self, titulo, itens, vazio_label='Nenhum registro'):
        self.story.append(Paragraph(titulo, self.styles['SectionTitle']))
        if not itens:
            self.story.append(Paragraph(vazio_label, self.styles['Normal']))
            return

        texto = '<br/>'.join(itens)
        self.story.append(Paragraph(texto, self.styles['Normal']))

    def _format_endereco(self, endereco):
        numero = endereco.numero or 'S/N'
        complemento = f", {endereco.complemento}" if endereco.complemento else ""
        pais_part = f", {endereco.pais}" if endereco.pais else ""
        pais_normalizado = (endereco.pais or '').strip().lower()
        mostrar_cep = pais_normalizado in ['', 'brasil']
        cep_part = f" - CEP: {endereco.cep}" if mostrar_cep and endereco.cep else ""
        uf_part = f"/{endereco.uf}" if mostrar_cep and endereco.uf else ""
        return (
            f"{endereco.logradouro}, {numero}{complemento} - "
            f"{endereco.bairro}, {endereco.cidade}{uf_part}{pais_part}{cep_part}"
        )

    def adicionar_conteudo(self):
        fornecedor = self.fornecedor

        self._add_kv_table('Documentacao', [
            ('CNPJ/CPF', fornecedor.get_cnpj_cpf_formatado() if hasattr(fornecedor, 'get_cnpj_cpf_formatado') else self._fmt(fornecedor.cnpj_cpf)),
            ('Inscricao Estadual', self._fmt(fornecedor.inscricao_estadual)),
            ('Inscricao Municipal', self._fmt(fornecedor.inscricao_municipal)),
            ('Codigo EORI', self._fmt(fornecedor.eori)),
            ('Regime Tributario', self._fmt(getattr(fornecedor, 'get_regime_tributario_display', None)() if hasattr(fornecedor, 'get_regime_tributario_display') else fornecedor.regime_tributario)),
        ])

        self._add_kv_table('Contato', [
            ('Telefone', self._fmt(fornecedor.telefone)),
            ('Celular/WhatsApp', self._fmt(fornecedor.celular)),
            ('Email', self._fmt(fornecedor.email)),
            ('Website', self._fmt(fornecedor.website)),
        ])

        self._add_kv_table('Condicoes Comerciais', [
            ('Prazo de Pagamento', f"{fornecedor.prazo_pagamento_padrao} dias" if fornecedor.prazo_pagamento_padrao else '—'),
            ('Prazo de Entrega', f"{fornecedor.prazo_entrega_medio} dias" if fornecedor.prazo_entrega_medio else '—'),
            ('Moeda Padrao', self._fmt(fornecedor.moeda_padrao)),
            ('Porto Preferencial', self._fmt(fornecedor.porto_preferencial)),
        ])

        self._add_kv_table('Homologacao', [
            ('Data Homologacao', self._fmt(fornecedor.data_homologacao.strftime('%d/%m/%Y') if fornecedor.data_homologacao else None)),
            ('Validade Homologacao', self._fmt(fornecedor.validade_homologacao.strftime('%d/%m/%Y') if fornecedor.validade_homologacao else None)),
            ('Classificacao de Risco', self._fmt(getattr(fornecedor, 'get_classificacao_risco_display', None)() if hasattr(fornecedor, 'get_classificacao_risco_display') else fornecedor.classificacao_risco)),
        ])

        if fornecedor.observacoes:
            self._add_list('Observacoes', [self._fmt(fornecedor.observacoes)], vazio_label='—')

        enderecos_linhas = []
        for endereco in fornecedor.enderecos.all():
            label = endereco.get_tipo_display()
            desc = self._format_endereco(endereco)
            if endereco.padrao:
                label = f"{label} (Principal)"
            enderecos_linhas.append((label, desc))

        if enderecos_linhas:
            self._add_kv_table('Enderecos', enderecos_linhas)
        else:
            self._add_list('Enderecos', [], vazio_label='Nenhum endereco cadastrado')

        contatos_itens = []
        for contato in fornecedor.contatos.all():
            partes = [contato.nome]
            if contato.cargo:
                partes.append(contato.cargo)
            if contato.telefone:
                partes.append(contato.telefone)
            if contato.email:
                partes.append(contato.email)
            contatos_itens.append(' | '.join(partes))

        self._add_list('Contatos', contatos_itens, vazio_label='Nenhum contato cadastrado')

        contas_itens = []
        for conta in fornecedor.contas_bancarias.all():
            partes = [conta.banco]
            if conta.agencia:
                partes.append(f"Agencia {conta.agencia}{('-' + conta.agencia_digito) if conta.agencia_digito else ''}")
            if conta.conta:
                partes.append(f"Conta {conta.conta}{('-' + conta.conta_digito) if conta.conta_digito else ''}")
            if conta.tipo_conta:
                partes.append(conta.get_tipo_conta_display())
            contas_itens.append(' | '.join(partes))

        self._add_list('Contas Bancarias', contas_itens, vazio_label='Nenhuma conta cadastrada')

    def gerar(self):
        self.doc.build(self.story)
        self.buffer.seek(0)
        return self.buffer.getvalue()


class RelatorioFornecedoresExcel:
    def __init__(self, empresa, titulo='Relatorio de Fornecedores'):
        self.empresa = empresa
        self.titulo = titulo
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = 'Fornecedores'
        self._setup_styles()

    def _setup_styles(self):
        self.header_fill = PatternFill(start_color='1f4788', end_color='1f4788', fill_type='solid')
        self.header_font = Font(bold=True, color='FFFFFF', size=11)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def adicionar_cabecalho(self, filtros_desc=''):
        self.ws['A1'] = self.titulo
        self.ws['A1'].font = Font(bold=True, size=14, color='1f4788')

        self.ws['A2'] = f'Empresa: {self.empresa.nome}'
        self.ws['A3'] = f'Gerado em: {datetime.now().strftime("%d/%m/%Y as %H:%M")}'

        if filtros_desc:
            self.ws['A4'] = f'Filtros: {filtros_desc}'

        return 5

    def adicionar_tabela(self, fornecedores, linha_inicio=5):
        colunas = ['Razao Social', 'CNPJ/CPF', 'Cidade/UF', 'Status']

        for col, titulo in enumerate(colunas, 1):
            cell = self.ws.cell(row=linha_inicio, column=col)
            cell.value = titulo
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.border = self.border
            cell.alignment = Alignment(horizontal='center', vertical='center')

        linha = linha_inicio + 1
        for fornecedor in fornecedores:
            cnpj_cpf = fornecedor.get_cnpj_cpf_formatado() if hasattr(fornecedor, 'get_cnpj_cpf_formatado') else (fornecedor.cnpj_cpf or '')
            if fornecedor.cidade and fornecedor.uf:
                cidade_uf = f"{fornecedor.cidade}/{fornecedor.uf}"
            elif fornecedor.cidade:
                cidade_uf = fornecedor.cidade
            elif fornecedor.uf:
                cidade_uf = fornecedor.uf
            else:
                cidade_uf = ''

            status = 'Ativo' if fornecedor.ativa else 'Inativo'
            self.ws.cell(row=linha, column=1).value = fornecedor.nome
            self.ws.cell(row=linha, column=2).value = cnpj_cpf
            self.ws.cell(row=linha, column=3).value = cidade_uf
            self.ws.cell(row=linha, column=4).value = status

            for col in range(1, 5):
                cell = self.ws.cell(row=linha, column=col)
                cell.border = self.border
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

            linha += 1

        self.ws.column_dimensions['A'].width = 40
        self.ws.column_dimensions['B'].width = 18
        self.ws.column_dimensions['C'].width = 18
        self.ws.column_dimensions['D'].width = 12

        return linha

    def gerar(self):
        buffer = io.BytesIO()
        self.wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
